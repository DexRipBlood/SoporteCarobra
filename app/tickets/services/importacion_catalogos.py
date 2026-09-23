import hashlib
import re
import secrets
import unicodedata
from decimal import Decimal, InvalidOperation

from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from tickets.models import AsignacionPersonal, Empresa, GrupoTrabajo, MembresiaGrupo, Persona, Tienda, Zona
from usuarios.models import PerfilUsuario


User = get_user_model()


ENCABEZADOS_TIENDA = {
    "ID": "ID_DE_TIENDA",
    "ID_TIENDA": "ID_DE_TIENDA",
    "ENCARGADO": "ENCARGADO_DE_SOPORTE",
    "NOMBRE_TIENDA": "NOMBRE_DE_LA_TIENDA",
    # Compatibilidad con la plantilla anterior.
    "CODIGO_TIENDA": "ID_DE_TIENDA",
    "NUMERO_TIENDA": "TIENDA",
    "NOMBRE_ZONA": "DISTRITAL",
    "MUNICIPIO": "CIUDAD",
    "LOCALIDAD": "CIUDAD",
}
ENCABEZADOS_DIRECTORIO = {
    "NUMERO_DE_EMPLEADO": "ID_PERSONAL", "TELEFONO_CEL": "TELEFONO",
    "CORREO_ELECTRONICO": "CORREO", "KEY": "KEY_TIENDA",
    "NUMERODETIENDA": "NUMERO_TIENDA", "STATUS": "ACTIVO",
}

REQUERIDAS_TIENDA = {
    "ID_DE_TIENDA",
    "TIENDA",
    "NOMBRE_DE_LA_TIENDA",
    "ESTADO",
    "DISTRITAL",
}
REQUERIDAS_COORDENADAS = {"ID_DE_TIENDA", "LATITUD", "LONGITUD"}

VALORES_SI = {"SI", "S", "1", "TRUE", "VERDADERO", "ACTIVO", "ACTIVA", "YES"}
VALORES_NO = {"NO", "N", "0", "FALSE", "FALSO", "INACTIVO", "INACTIVA", ""}


def _clave(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).upper().strip()
    return re.sub(r"[^A-Z0-9]+", "_", texto).strip("_")


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _ciudad_desde_direccion(direccion):
    """Recupera la localidad escrita después del código postal, si existe."""
    coincidencia = re.search(r"\b\d{5}\s*,?\s+([^,]+)", str(direccion or ""), re.IGNORECASE)
    if not coincidencia:
        return ""
    return coincidencia.group(1).strip(" .;-")[:140]


def _booleano(valor, defecto=True):
    texto = _clave(valor)
    if not texto:
        return defecto
    return texto not in {"NO", "N", "0", "FALSE", "INACTIVO", "INACTIVA"}


def _booleano_catalogo(valor, defecto=False):
    clave = _clave(valor)
    if not clave:
        return defecto
    if clave in VALORES_SI:
        return True
    if clave in VALORES_NO:
        return False
    raise ValueError(f"valor Sí/No no reconocido: {valor}")


def _encabezados_normalizados(encabezados_originales, tipo):
    encabezados = []
    posiciones = {}
    for posicion, valor in enumerate(encabezados_originales, start=1):
        clave = _clave(valor)
        if tipo in {"TIENDAS", "COORDENADAS"}:
            clave = ENCABEZADOS_TIENDA.get(clave, clave)
        elif tipo == "DIRECTORIO":
            clave = ENCABEZADOS_DIRECTORIO.get(clave, clave)
        if clave and clave in posiciones:
            raise ValueError(
                f"La columna {clave.replace('_', ' ')} está repetida "
                f"(posiciones {posiciones[clave]} y {posicion})."
            )
        if clave:
            posiciones[clave] = posicion
        encabezados.append(clave)
    return encabezados


def _codigo_zona(empresa, estado, distrital, codigo_anterior=""):
    if codigo_anterior:
        return codigo_anterior[:60]
    identidad = f"{empresa.pk}|{_clave(estado)}|{_clave(distrital)}".encode()
    resumen = hashlib.sha1(identidad).hexdigest()[:10].upper()
    return f"DIST-{resumen}"


def _indice_personas():
    indice = {}
    for persona in Persona.objects.filter(activa=True):
        indice.setdefault(_clave(persona.nombre), []).append(persona)
    return indice


def _persona_unica(indice, nombre):
    coincidencias = indice.get(_clave(nombre), []) if nombre else []
    return coincidencias[0] if len(coincidencias) == 1 else None


def _coordenadas(fila, obligatorias=False):
    latitud_texto = fila.get("LATITUD", "")
    longitud_texto = fila.get("LONGITUD", "")
    if not latitud_texto and not longitud_texto and not obligatorias:
        return None
    if not latitud_texto or not longitud_texto:
        raise ValueError("LATITUD y LONGITUD deben capturarse juntas")
    try:
        latitud = Decimal(latitud_texto.replace(",", "."))
        longitud = Decimal(longitud_texto.replace(",", "."))
    except (InvalidOperation, AttributeError):
        raise ValueError("LATITUD y LONGITUD deben ser números decimales")
    if not (Decimal("14") <= latitud <= Decimal("33.5")):
        raise ValueError("LATITUD está fuera del rango de México")
    if not (Decimal("-119") <= longitud <= Decimal("-86")):
        raise ValueError("LONGITUD está fuera del rango de México")
    return latitud, longitud


def analizar_catalogo(archivo, tipo, empresa=None):
    libro = load_workbook(archivo, read_only=True, data_only=True)
    hoja = libro.active
    filas = hoja.iter_rows(values_only=True)
    encabezados_originales = next(filas, None)
    if not encabezados_originales:
        raise ValueError("El archivo no contiene encabezados.")
    encabezados = _encabezados_normalizados(encabezados_originales, tipo)
    requeridas = {
        "TIENDAS": REQUERIDAS_TIENDA,
        "COORDENADAS": REQUERIDAS_COORDENADAS,
        "USUARIOS": {"CORREO", "NOMBRE"},
        "DIRECTORIO": {"ID_PERSONAL", "NOMBRE"},
    }[tipo]
    faltantes = sorted(requeridas - set(encabezados))
    if faltantes:
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(faltantes))

    datos, errores = [], []
    ids_archivo = set()
    for numero, valores in enumerate(filas, start=2):
        fila = {
            encabezados[i]: _texto(valor)
            for i, valor in enumerate(valores)
            if i < len(encabezados) and encabezados[i]
        }
        if not any(fila.values()):
            continue
        problemas = []
        for campo in requeridas:
            if not fila.get(campo):
                problemas.append(f"{campo} está vacío")
        if tipo in {"USUARIOS", "DIRECTORIO"} and fila.get("CORREO") and "@" not in fila["CORREO"]:
            problemas.append("CORREO no es válido")
        if tipo in {"TIENDAS", "COORDENADAS"}:
            fila["ID_DE_TIENDA"] = fila.get("ID_DE_TIENDA", "").upper()
            id_tienda = _clave(fila.get("ID_DE_TIENDA"))
            if id_tienda in ids_archivo:
                problemas.append("ID DE TIENDA está repetido en el archivo")
            elif id_tienda:
                ids_archivo.add(id_tienda)
            try:
                _coordenadas(fila, obligatorias=tipo == "COORDENADAS")
            except ValueError as error:
                problemas.append(str(error))
            if tipo == "TIENDAS":
                for campo in ("RETIRO", "PILOTO"):
                    try:
                        _booleano_catalogo(fila.get(campo), defecto=False)
                    except ValueError as error:
                        problemas.append(f"{campo}: {error}")
        if problemas:
            errores.append({"fila": numero, "errores": problemas})
        else:
            fila["_FILA"] = numero
            datos.append(fila)
    resultado = {
        "tipo": tipo,
        "filas": datos,
        "errores": errores,
        "total": len(datos) + len(errores),
    }
    if tipo in {"TIENDAS", "COORDENADAS"} and empresa is not None:
        ids_existentes = set(
            Tienda.objects.filter(
                empresa=empresa,
                id_externo__in=[fila["ID_DE_TIENDA"] for fila in datos],
            ).values_list("id_externo", flat=True)
        )
        if tipo == "COORDENADAS":
            filas_validas = []
            for fila in datos:
                if fila["ID_DE_TIENDA"] not in ids_existentes:
                    errores.append({
                        "fila": fila["_FILA"],
                        "errores": ["ID DE TIENDA no existe en la empresa seleccionada"],
                    })
                else:
                    filas_validas.append(fila)
            resultado["filas"] = filas_validas
            resultado["errores"] = errores
            resultado["resumen"] = {"nuevas": 0, "actualizables": len(filas_validas)}
        else:
            resultado["resumen"] = {
                "nuevas": sum(fila["ID_DE_TIENDA"] not in ids_existentes for fila in datos),
                "actualizables": sum(fila["ID_DE_TIENDA"] in ids_existentes for fila in datos),
            }
    return resultado


def _username_disponible(correo, solicitado=""):
    base = _clave(solicitado or correo.split("@", 1)[0]).lower().replace("_", ".")[:140] or "usuario"
    candidato, numero = base, 1
    while User.objects.filter(username__iexact=candidato).exists():
        numero += 1
        candidato = f"{base[:140-len(str(numero))]}{numero}"
    return candidato


@transaction.atomic
def importar_catalogo(previsualizacion, empresa=None):
    nuevos = actualizados = sin_cambios = 0
    usuarios_nuevos = []
    if previsualizacion["tipo"] == "TIENDAS":
        if empresa is None:
            raise ValueError("La empresa es obligatoria.")
        personas = _indice_personas()
        for fila in previsualizacion["filas"]:
            distrital = fila["DISTRITAL"]
            zona = Zona.objects.filter(
                empresa=empresa,
                nombre__iexact=distrital,
                activa=True,
            ).first()
            datos_zona = {
                "nombre": distrital,
                "estado": fila["ESTADO"],
                "distrital_nombre": fila.get("GERENTE_DISTRITAL", ""),
                "activa": True,
            }
            if zona is None:
                zona, _ = Zona.objects.update_or_create(
                    empresa=empresa,
                    codigo=_codigo_zona(empresa, fila["ESTADO"], distrital, fila.get("ID_ZONA", "")),
                    defaults=datos_zona,
                )
            else:
                for campo, valor in datos_zona.items():
                    setattr(zona, campo, valor)
                zona.save()
            retiro = _booleano_catalogo(fila.get("RETIRO"), defecto=False)
            direccion = fila.get("DIRECCION", "")
            ciudad = fila.get("CIUDAD", "") or _ciudad_desde_direccion(direccion)
            coordenadas = _coordenadas(fila)
            defaults = {
                "zona": zona,
                "clave": fila.get("KEY", ""),
                "numero": fila["TIENDA"],
                "nombre": fila["NOMBRE_DE_LA_TIENDA"],
                "estado": fila["ESTADO"],
                "ciudad": ciudad,
                "direccion": direccion,
                "agencia": fila.get("AGENCIA", ""),
                "socio": fila.get("SOCIO", ""),
                "cadena": Tienda.codigo_tipo_desde_socio(fila.get("SOCIO", "")),
                "gerente_distrital": fila.get("GERENTE_DISTRITAL", ""),
                "coordinador": fila.get("COORDINADOR", ""),
                "tipo_coordinador": fila.get("TIPO_DE_COORDINADOR", fila.get("TIPO_COORDINADOR", "")),
                "encargado_soporte_nombre": fila.get("ENCARGADO_DE_SOPORTE", ""),
                "partner_nombre": fila.get("PARTNER", ""),
                "lider_actual_nombre": fila.get("LIDER_ACTUAL", ""),
                "encargado_soporte": _persona_unica(personas, fila.get("ENCARGADO_DE_SOPORTE")),
                "partner": _persona_unica(personas, fila.get("PARTNER")),
                "lider": _persona_unica(personas, fila.get("LIDER_ACTUAL")),
                "retiro": retiro,
                "prioridad_operativa": fila.get("PRIORIDAD", ""),
                "piloto": _booleano_catalogo(fila.get("PILOTO"), defecto=False),
                "activa": not retiro,
            }
            if coordenadas:
                defaults.update({
                    "latitud": coordenadas[0],
                    "longitud": coordenadas[1],
                    "estado_geocodificacion": Tienda.EstadoGeocodificacion.LOCALIZADA,
                    "detalle_geocodificacion": "Coordenadas proporcionadas en la importación.",
                    "geocodificada_at": timezone.now(),
                })
            tienda = Tienda.objects.filter(
                empresa=empresa,
                id_externo=fila["ID_DE_TIENDA"],
            ).first()
            if tienda is None:
                legado = Tienda.objects.filter(
                    empresa=empresa,
                    id_externo="",
                    numero=fila["TIENDA"],
                )
                if legado.count() == 1:
                    tienda = legado.first()
                    tienda.id_externo = fila["ID_DE_TIENDA"]
                    tienda.save(update_fields=["id_externo", "actualizado_at"])
            creada = tienda is None
            if creada:
                tienda = Tienda.objects.create(
                    empresa=empresa,
                    id_externo=fila["ID_DE_TIENDA"],
                    **defaults,
                )
            if creada:
                nuevos += 1
            else:
                if tienda.direccion != direccion and not coordenadas:
                    defaults.update({
                        "latitud": None,
                        "longitud": None,
                        "estado_geocodificacion": Tienda.EstadoGeocodificacion.PENDIENTE,
                        "detalle_geocodificacion": "Dirección actualizada; pendiente de localizar.",
                        "geocodificada_at": None,
                    })
                cambios = any(getattr(tienda, campo) != valor for campo, valor in defaults.items())
                if cambios:
                    for campo, valor in defaults.items():
                        setattr(tienda, campo, valor)
                    tienda.save()
                    actualizados += 1
                else:
                    sin_cambios += 1
    elif previsualizacion["tipo"] == "COORDENADAS":
        if empresa is None:
            raise ValueError("La empresa es obligatoria.")
        for fila in previsualizacion["filas"]:
            tienda = Tienda.objects.get(
                empresa=empresa,
                id_externo=fila["ID_DE_TIENDA"],
            )
            latitud, longitud = _coordenadas(fila, obligatorias=True)
            cambios = tienda.latitud != latitud or tienda.longitud != longitud
            tienda.latitud = latitud
            tienda.longitud = longitud
            tienda.estado_geocodificacion = Tienda.EstadoGeocodificacion.LOCALIZADA
            tienda.detalle_geocodificacion = "Coordenadas actualizadas mediante Excel."
            tienda.geocodificada_at = timezone.now()
            tienda.save(update_fields=[
                "latitud", "longitud", "estado_geocodificacion",
                "detalle_geocodificacion", "geocodificada_at", "actualizado_at",
            ])
            if cambios:
                actualizados += 1
            else:
                sin_cambios += 1
    elif previsualizacion["tipo"] == "USUARIOS":
        for fila in previsualizacion["filas"]:
            correo = fila["CORREO"].lower()
            usuario = User.objects.filter(email__iexact=correo).first()
            creada = usuario is None
            if creada:
                usuario = User(username=_username_disponible(correo, fila.get("USUARIO", "")), email=correo)
                # Clave aleatoria no comunicada: el usuario define la suya
                # mediante la invitación de un solo uso o entra con Google.
                usuario.set_password(secrets.token_urlsafe(48))
            nombre_anterior = (usuario.first_name, usuario.last_name, usuario.is_active)
            usuario.first_name = fila["NOMBRE"]
            usuario.last_name = fila.get("APELLIDOS", "")
            usuario.is_active = _booleano(fila.get("ACTIVO"))
            usuario.save()
            perfil, _ = PerfilUsuario.objects.get_or_create(user=usuario)
            perfil.numero_empleado = fila.get("NUMERO_EMPLEADO", perfil.numero_empleado)
            perfil.telefono = fila.get("TELEFONO", perfil.telefono)
            perfil.area = fila.get("AREA", perfil.area)
            perfil.cargo = fila.get("PUESTO", fila.get("CARGO", perfil.cargo))
            perfil.rol = PerfilUsuario.Rol.ADMIN if _clave(fila.get("ROL")) in {"ADMIN", "ADMINISTRADOR"} else PerfilUsuario.Rol.USUARIO
            perfil.activo = usuario.is_active
            perfil.save()
            persona_id = fila.get("ID_PERSONAL") or fila.get("NUMERO_EMPLEADO")
            if persona_id:
                Persona.objects.update_or_create(id_personal=persona_id, defaults={"nombre": usuario.get_full_name() or usuario.username, "puesto": perfil.cargo, "horario": fila.get("HORARIO", ""), "telefono": perfil.telefono, "correo": correo, "usuario": usuario, "activa": usuario.is_active})
            empresa_codigo, grupo_nombre = fila.get("ID_EMPRESA"), fila.get("GRUPO")
            if empresa_codigo and grupo_nombre:
                empresa_usuario = Empresa.objects.filter(codigo__iexact=empresa_codigo).first()
                if empresa_usuario:
                    grupo, _ = GrupoTrabajo.objects.get_or_create(empresa=empresa_usuario, nombre=grupo_nombre)
                    nivel = _clave(fila.get("NIVEL"))
                    if nivel not in {"N1", "N2", "N3"}:
                        nivel = "N1"
                    MembresiaGrupo.objects.update_or_create(grupo=grupo, usuario=usuario, defaults={"nivel": nivel, "activa": usuario.is_active})
            if creada:
                nuevos += 1
                usuarios_nuevos.append(usuario)
            elif nombre_anterior != (usuario.first_name, usuario.last_name, usuario.is_active):
                actualizados += 1
            else:
                sin_cambios += 1
    else:  # Directorio de Activos: no crea cuentas, solo contactos operativos.
        for fila in previsualizacion["filas"]:
            persona, creada = Persona.objects.update_or_create(
                id_personal=fila["ID_PERSONAL"],
                defaults={"nombre": fila["NOMBRE"], "puesto": fila.get("PUESTO", ""), "horario": fila.get("HORARIO", ""), "descanso": fila.get("DESCANSO", ""), "telefono": fila.get("TELEFONO", ""), "correo": fila.get("CORREO", ""), "activa": _booleano(fila.get("ACTIVO"))},
            )
            if creada: nuevos += 1
            else: actualizados += 1
            tiendas = Tienda.objects.filter(activa=True)
            if fila.get("KEY_TIENDA"):
                tiendas = tiendas.filter(clave__iexact=fila["KEY_TIENDA"])
            elif fila.get("NUMERO_TIENDA"):
                tiendas = tiendas.filter(numero__iexact=fila["NUMERO_TIENDA"])
            if tiendas.count() == 1:
                tienda = tiendas.first()
                AsignacionPersonal.objects.update_or_create(
                    persona=persona, empresa=tienda.empresa, tienda=tienda,
                    defaults={"alcance": AsignacionPersonal.Alcance.TIENDA, "funcion": AsignacionPersonal.Funcion.PERSONAL, "horario": persona.horario, "activa": persona.activa},
                )
    return {"nuevos": nuevos, "actualizados": actualizados, "sin_cambios": sin_cambios, "usuarios_nuevos": usuarios_nuevos}

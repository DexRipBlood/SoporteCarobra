from collections import Counter
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


# =========================================================
# PLANTILLA BRADESCARD
# =========================================================

COLUMNAS_ESPERADAS = {
    "tienda": "Tienda",
    "ticket_bradescard": "Ticket Bradescard",
    "estatus": "Estatus",
    "fecha_apertura": "Fecha de Apertura",
    "dias_transcurridos": "Dias Transcurridos",
    "categoria": "Categoria",
    "incidencia_general": "Incidencia General",
    "incidencia_especifica": "Incidencia Especifica",
    "seguimientos_descripcion": "Seguimientos - Descripcion",
    "fecha_reporte": "Fecha del Reporte",
    "hora_reporte": "Hora Reporte",
    "ultima_modificacion": "Ultima Modificacion",
    "dias_sin_comentar": "Dias Sin Comentar",
    "tiempo_transcurrido": "Tiempo Transcurrido",
    "asignado": "Asignado",
    "type": "Type",
    "usuario": "Usuario",
    "tiempo_total": "Tiempo Total",
    "soc": "SOC",
    "tablet": "Tablet",
    "cierre_por_falta": "Cierre por Falta",
}


def normalizar_texto(valor):
    """
    Convierte encabezados a una forma comparable.

    Ejemplos:
        Categoría -> categoria
        Última Modificación -> ultima modificacion
        Seguimientos - Descripción -> seguimientos descripcion
    """

    if valor is None:
        return ""

    texto = str(valor).strip()

    texto = unicodedata.normalize("NFKD", texto)

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(caracter)
    )

    texto = texto.lower()

    texto = re.sub(
        r"[^a-z0-9]+",
        " ",
        texto,
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return texto.strip()


COLUMNAS_NORMALIZADAS = {
    normalizar_texto(nombre): clave
    for clave, nombre in COLUMNAS_ESPERADAS.items()
}


def normalizar_ticket(valor):
    """
    Evita tickets como:
        357347.0

    cuando Excel los entrega como número.
    """

    if valor is None:
        return ""

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    return str(valor).strip()


def buscar_fila_encabezados(ws):
    """
    Busca los encabezados en las primeras 10 filas.
    Esto permite tolerar una fila de título si alguna vez
    modifican ligeramente el archivo.
    """

    limite = min(
        ws.max_row or 10,
        10,
    )

    for numero_fila, fila in enumerate(
        ws.iter_rows(
            min_row=1,
            max_row=limite,
            values_only=True,
        ),
        start=1,
    ):

        valores = {
            normalizar_texto(valor)
            for valor in fila
            if valor is not None
        }

        if (
            "tienda" in valores
            and "ticket bradescard" in valores
        ):
            return numero_fila

    return None


def validar_excel_bradescard(ruta_archivo):
    """
    Lee y valida la plantilla.

    IMPORTANTE:
    Esta función NO guarda nada en la base de datos.
    """

    ruta = Path(ruta_archivo)

    resultado = {
        "valido": False,
        "archivo": ruta.name,
        "hoja": None,
        "fila_encabezados": None,
        "columnas_encontradas": 0,
        "columnas_reconocidas": 0,
        "columnas_faltantes": [],
        "columnas_extra": [],
        "total_filas": 0,
        "tickets_validos": 0,
        "tickets_sin_numero": 0,
        "filas_sin_ticket": [],
        "tickets_duplicados": 0,
        "duplicados": [],
        "errores": [],
    }

    # -----------------------------------------------------
    # Validaciones básicas
    # -----------------------------------------------------

    if not ruta.exists():

        resultado["errores"].append(
            "El archivo no existe."
        )

        return resultado

    if ruta.suffix.lower() != ".xlsx":

        resultado["errores"].append(
            "El archivo debe tener extensión .xlsx."
        )

        return resultado

    # -----------------------------------------------------
    # Abrir Excel
    # -----------------------------------------------------

    try:

        workbook = load_workbook(
            filename=ruta,
            read_only=True,
            data_only=True,
        )

    except Exception as error:

        resultado["errores"].append(
            f"No se pudo abrir el archivo: {error}"
        )

        return resultado

    try:

        ws = workbook.active

        resultado["hoja"] = ws.title

        # -------------------------------------------------
        # Buscar encabezados
        # -------------------------------------------------

        fila_encabezados = buscar_fila_encabezados(ws)

        if fila_encabezados is None:

            resultado["errores"].append(
                "No se encontró la fila de encabezados "
                "de la plantilla Bradescard."
            )

            return resultado

        resultado["fila_encabezados"] = fila_encabezados

        encabezados = next(
            ws.iter_rows(
                min_row=fila_encabezados,
                max_row=fila_encabezados,
                values_only=True,
            )
        )

        columnas = {}

        columnas_extra = []

        for indice, encabezado in enumerate(encabezados):

            if encabezado is None:
                continue

            normalizado = normalizar_texto(encabezado)

            if normalizado in COLUMNAS_NORMALIZADAS:

                clave = COLUMNAS_NORMALIZADAS[
                    normalizado
                ]

                columnas[clave] = indice

            else:

                columnas_extra.append(
                    str(encabezado).strip()
                )

        # -------------------------------------------------
        # Columnas
        # -------------------------------------------------

        resultado["columnas_encontradas"] = len(
            [
                x
                for x in encabezados
                if x is not None
            ]
        )

        resultado["columnas_reconocidas"] = len(
            columnas
        )

        faltantes = []

        for clave, nombre in COLUMNAS_ESPERADAS.items():

            if clave not in columnas:
                faltantes.append(nombre)

        resultado["columnas_faltantes"] = faltantes
        resultado["columnas_extra"] = columnas_extra

        if faltantes:

            resultado["errores"].append(
                "La plantilla no contiene todas "
                "las columnas requeridas."
            )

            return resultado

        # -------------------------------------------------
        # Validación de tickets
        # -------------------------------------------------

        indice_ticket = columnas[
            "ticket_bradescard"
        ]

        tickets = []

        filas_sin_ticket = []

        for numero_fila, fila in enumerate(
            ws.iter_rows(
                min_row=fila_encabezados + 1,
                values_only=True,
            ),
            start=fila_encabezados + 1,
        ):

            # Ignorar filas completamente vacías
            if not any(
                valor not in (None, "")
                for valor in fila
            ):
                continue

            resultado["total_filas"] += 1

            if indice_ticket >= len(fila):
                filas_sin_ticket.append(numero_fila)
                continue

            ticket = normalizar_ticket(
                fila[indice_ticket]
            )

            if not ticket:

                filas_sin_ticket.append(numero_fila)
                continue

            tickets.append(ticket)

        resultado["tickets_validos"] = len(
            tickets
        )

        resultado["tickets_sin_numero"] = len(
            filas_sin_ticket
        )

        resultado["filas_sin_ticket"] = (
            filas_sin_ticket[:20]
        )

        # -------------------------------------------------
        # Duplicados
        # -------------------------------------------------

        contador = Counter(tickets)

        duplicados = [
            ticket
            for ticket, cantidad in contador.items()
            if cantidad > 1
        ]

        resultado["tickets_duplicados"] = len(
            duplicados
        )

        resultado["duplicados"] = duplicados[:20]

        # -------------------------------------------------
        # Resultado final
        # -------------------------------------------------

        if filas_sin_ticket:

            resultado["errores"].append(
                "Existen filas sin Ticket Bradescard."
            )

        if duplicados:

            resultado["errores"].append(
                "Existen números de ticket duplicados "
                "dentro del mismo archivo."
            )

        resultado["valido"] = (
            not resultado["errores"]
        )

        return resultado

    finally:

        workbook.close()


    # =========================================================
# LECTOR DE FILAS PARA IMPORTACIÓN
# =========================================================

def obtener_filas_bradescard(ruta_archivo):
    """
    Convierte el Excel validado en una lista de diccionarios.

    Esta función todavía NO guarda nada en PostgreSQL.
    """

    ruta = Path(ruta_archivo)

    validacion = validar_excel_bradescard(ruta)

    if not validacion["valido"]:
        raise ValueError(
            "El archivo no pasó la validación."
        )

    workbook = load_workbook(
        filename=ruta,
        read_only=True,
        data_only=True,
    )

    try:
        ws = workbook.active

        fila_encabezados = validacion[
            "fila_encabezados"
        ]

        encabezados = next(
            ws.iter_rows(
                min_row=fila_encabezados,
                max_row=fila_encabezados,
                values_only=True,
            )
        )

        columnas = {}

        for indice, encabezado in enumerate(encabezados):

            if encabezado is None:
                continue

            normalizado = normalizar_texto(
                encabezado
            )

            if normalizado in COLUMNAS_NORMALIZADAS:

                clave = COLUMNAS_NORMALIZADAS[
                    normalizado
                ]

                columnas[clave] = indice

        filas = []

        for numero_fila, fila in enumerate(
            ws.iter_rows(
                min_row=fila_encabezados + 1,
                values_only=True,
            ),
            start=fila_encabezados + 1,
        ):

            if not any(
                valor not in (None, "")
                for valor in fila
            ):
                continue

            datos = {
                clave: (
                    fila[indice]
                    if indice < len(fila)
                    else None
                )
                for clave, indice in columnas.items()
            }

            datos["ticket_bradescard"] = (
                normalizar_ticket(
                    datos.get(
                        "ticket_bradescard"
                    )
                )
            )

            datos["_fila_excel"] = numero_fila

            filas.append(datos)

        return filas

    finally:
        workbook.close()


    # =========================================================
# HASH / COMPARACIÓN CONTRA POSTGRESQL
# =========================================================

import hashlib
import json
from datetime import date, datetime, time
from decimal import Decimal


def valor_serializable(valor):
    """
    Convierte valores provenientes de Excel a una
    representación estable para calcular hashes.

    Importante:
    26 y 26.0 deben producir exactamente el mismo valor.
    """

    if valor is None:
        return None

    if isinstance(valor, datetime):
        return valor.isoformat()

    if isinstance(valor, date):
        return valor.isoformat()

    if isinstance(valor, time):
        return valor.isoformat()

    # bool debe comprobarse antes porque bool hereda de int.
    if isinstance(valor, bool):
        return valor

    if isinstance(valor, Decimal):
        if valor == valor.to_integral_value():
            return int(valor)

        return str(valor.normalize())

    if isinstance(valor, float):
        # 26.0 -> 26
        if valor.is_integer():
            return int(valor)

        return valor

    return valor


def calcular_hash_fila(datos):
    """
    Genera un SHA-256 del contenido importado.

    Nos permitirá saber si un ticket existente
    realmente cambió respecto al Excel anterior.
    """

    datos_hash = {
        clave: valor_serializable(valor)
        for clave, valor in datos.items()
        if clave != "_fila_excel"
    }

    contenido = json.dumps(
        datos_hash,
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
        default=str,
    )

    return hashlib.sha256(
        contenido.encode("utf-8")
    ).hexdigest()


def analizar_importacion_bradescard(ruta_archivo, empresa=None, categorias=None):
    """
    PREVISUALIZACIÓN / DRY RUN.

    Consulta PostgreSQL pero NO CREA,
    NO ACTUALIZA y NO ELIMINA registros.
    """

    from tickets.models import Ticket

    filas = obtener_filas_bradescard(
        ruta_archivo
    )
    from .filtros_importacion import categoria_fila, filtrar_categorias
    categorias_disponibles = sorted({categoria_fila(f) for f in filas})
    total_archivo = len(filas)
    filas = filtrar_categorias(filas, categorias)

    numeros_ticket = [
        fila["ticket_bradescard"]
        for fila in filas
    ]

    existentes = {
        ticket.ticket_bradescard: ticket
        for ticket in Ticket.objects.filter(
            ticket_bradescard__in=numeros_ticket,
            empresa=empresa,
        ).only(
            "id",
            "ticket_bradescard",
            "hash_origen",
        )
    }

    nuevos = []
    actualizados = []
    sin_cambios = []

    for fila in filas:

        numero = fila[
            "ticket_bradescard"
        ]

        hash_nuevo = calcular_hash_fila(
            fila
        )

        ticket_existente = existentes.get(
            numero
        )

        if ticket_existente is None:

            nuevos.append(numero)

            continue

        if (
            ticket_existente.hash_origen
            == hash_nuevo
        ):

            sin_cambios.append(numero)

        else:

            actualizados.append(numero)

    return {
        "total": len(filas),
        "omitidos": total_archivo - len(filas),
        "categorias_disponibles": categorias_disponibles,
        "categorias_seleccionadas": categorias_disponibles if categorias is None else categorias,
        "existentes_en_bd": len(existentes),
        "nuevos": len(nuevos),
        "actualizados": len(actualizados),
        "sin_cambios": len(sin_cambios),

        "muestra_nuevos": nuevos[:10],
        "muestra_actualizados": actualizados[:10],
        "muestra_sin_cambios": sin_cambios[:10],

        "se_guardo_algo": False,
    }

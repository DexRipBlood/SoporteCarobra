import mimetypes
import re
from datetime import datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from uuid import uuid4


from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Case, Count, IntegerField, Q, Value, When

from .models import (
    AsignacionPersonal,
    ArchivoTicket,
    ComentarioTicket,
    ConfiguracionSLA,
    Empresa,
    EstadoTicket,
    GrupoTrabajo,
    HistorialTicket,
    ImportacionTickets,
    MembresiaGrupo,
    OrigenTicket,
    Persona,
    PrioridadTicket,
    ProgramacionTrabajo,
    Ticket,
    Tienda,
    Zona,
)

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import FileResponse, Http404, HttpResponse, JsonResponse, StreamingHttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .forms import (
    ArchivoTicketForm,
    AsignacionPersonalForm,
    AsignarResponsableForm,
    ComentarioTicketForm,
    ConfiguracionSLAForm,
    CerrarTicketForm,
    CoordenadasTiendaForm,
    EmpresaForm,
    GestionTicketForm,
    GrupoTrabajoForm,
    ImportacionBradescardForm,
    ImportacionCatalogoForm,
    MembresiaGrupoForm,
    PersonaForm,
    ProgramacionTrabajoForm,
    ReabrirTicketForm,
    ReporteIncidenciasForm,
    TicketCarobraForm,
    TiendaForm,
    ZonaForm,
)
from .services.importador_excel import analizar_importacion_bradescard
from .services.importacion_bradescard import importar_bradescard
from .services.importacion_bradescard import (
    partner_automatico,
    responsable_automatico,
)
from .services.importacion_catalogos import analizar_catalogo, importar_catalogo
from .permisos import tickets_visibles_para, participa_en_ticket
from usuarios.permisos import administrador_required, es_administrador
from usuarios.correos import enviar_invitacion_usuario
from .services.operacion import resumen_sla, sincronizar_sla, seguimiento_actual, permiso_ticket, requieren_administracion, con_sla_actual, registrar_segmento_operativo
from .operacion_views import SeguimientoForm, SuplenciaForm


COLORES_ESTATUS_REPORTE = {
    "PARO TOTAL": "C62828",
    "INTERMITENCIA": "EF6C00",
    "MÓDEM DE CHIP": "1565C0",
    "MODEM DE CHIP": "1565C0",
    "RED ALTERNA": "2E7D32",
}

COLORES_SLA_REPORTE = {
    "EXCEDIDO": "C62828",
    "ATENCION": "1565C0",
    "PAUSA": "EF6C00",
}
PRIORIDAD_SLA_REPORTE = {"PAUSA": 1, "ATENCION": 2, "EXCEDIDO": 3}


def _nombre_reporte_usuario(usuario):
    return _nombre_usuario(usuario) if usuario else "Sin asignar"


def _segundos_en_rango(inicio, fin, desde, hasta):
    """Duración de una ventana, limitada al rango solicitado por el usuario."""
    inicio = max(inicio, desde)
    fin = min(fin or hasta, hasta)
    return max(0, int((fin - inicio).total_seconds()))


def _filas_reporte_incidencias(datos):
    """Obtiene tickets que se cruzan con el periodo y prepara filas auditables."""
    desde = timezone.make_aware(datetime.combine(datos["fecha_inicio"], time.min))
    hasta = timezone.make_aware(datetime.combine(datos["fecha_fin"], time.max))
    tickets = (
        Ticket.objects.filter(creado_at__lte=hasta)
        .filter(Q(cerrado_at__isnull=True) | Q(cerrado_at__gte=desde))
        .select_related(
            "tienda_registrada__zona", "responsable", "partner",
            "estatus_operativo_actual", "dependencia_actual",
        )
        .prefetch_related("segmentos_sla")
        .order_by("tienda", "creado_at")
    )
    if datos.get("tienda"):
        tickets = tickets.filter(tienda_registrada=datos["tienda"])
    if datos.get("coordinador"):
        tickets = tickets.filter(Q(responsable=datos["coordinador"]) | Q(tienda_registrada__zona__responsable=datos["coordinador"]))
    if datos.get("distrital"):
        tickets = tickets.filter(tienda_registrada__zona=datos["distrital"])
    if datos.get("incidencia"):
        texto = datos["incidencia"].strip()
        tickets = tickets.filter(Q(categoria__icontains=texto) | Q(incidencia_general__icontains=texto))

    filas = []
    for ticket in tickets:
        inicio = max(ticket.creado_at, desde)
        fin = min(ticket.cerrado_at or hasta, hasta)
        total = _segundos_en_rango(ticket.creado_at, ticket.cerrado_at, desde, hasta)
        efectivo = ticket.sla_legacy_segundos
        pausa = 0
        for segmento in ticket.segmentos_sla.all():
            segundos = _segundos_en_rango(segmento.inicio, segmento.fin, desde, hasta)
            if segmento.cuenta_sla:
                efectivo += segundos
            else:
                pausa += segundos
        estatus = ticket.estatus_operativo_actual.nombre if ticket.estatus_operativo_actual else "Sin definir"
        sla = resumen_sla(ticket)
        tienda = ticket.tienda_registrada
        zona = tienda.zona if tienda else None
        filas.append({
            "ticket": ticket,
            "tienda": tienda.nombre if tienda else ticket.tienda or "Sin tienda vinculada",
            "coordinador": _nombre_reporte_usuario(ticket.responsable or (zona.responsable if zona else None)),
            "distrital": zona.distrital_nombre or zona.nombre if zona else "—",
            "inicio": inicio,
            "fin": fin if ticket.cerrado_at else None,
            "total": total,
            "efectivo": efectivo,
            "pausa": pausa,
            "excedido": sla["excedido"],
            "dias_calendario": max(1, (fin.date() - inicio.date()).days + 1),
            "folios": ticket.ticket_bradescard or ticket.folio,
            "afectacion": ticket.incidencia_general or ticket.categoria or "Sin clasificar",
            "descripcion": ticket.incidencia_especifica or "—",
            "estatus": estatus,
            "dependencia": ticket.dependencia_actual.nombre if ticket.dependencia_actual else "—",
        })
    return filas


def _matriz_sla_reporte(filas, datos):
    """Agrupa por tienda y pinta cada día en que el ticket tuvo actividad SLA.

    La marca se genera desde SegmentoSLA, no desde una captura manual. Mientras
    no exista historial de estatus operativo por periodo, la matriz representa
    con precisión el tiempo efectivo y las pausas del SLA.
    """
    dias = []
    cursor = datos["fecha_inicio"]
    while cursor <= datos["fecha_fin"]:
        dias.append(cursor)
        cursor += timedelta(days=1)
    inicio_rango = timezone.make_aware(datetime.combine(datos["fecha_inicio"], time.min))
    fin_rango = timezone.make_aware(datetime.combine(datos["fecha_fin"], time.max))
    agrupadas = {}
    for fila in filas:
        clave = fila["tienda"]
        grupo = agrupadas.setdefault(clave, {
            "distrital": fila["distrital"], "coordinador": fila["coordinador"],
            "tienda": clave, "marcas": {}, "folios": [], "tipos": [],
        "comentarios": [], "efectivo": 0, "pausa": 0, "limites": [],
            "excedido": False,
        })
        ticket = fila["ticket"]
        grupo["folios"].append(fila["folios"])
        grupo["tipos"].append(fila["afectacion"])
        if fila["descripcion"] != "—":
            grupo["comentarios"].append(fila["descripcion"])
        grupo["efectivo"] += fila["efectivo"]
        grupo["pausa"] += fila["pausa"]
        if ticket.sla_limite_minutos is not None:
            grupo["limites"].append(ticket.sla_limite_minutos)
        grupo["excedido"] = grupo["excedido"] or fila["excedido"]
        segmentos = list(ticket.segmentos_sla.all())
        if not segmentos:
            segmentos = [type("Segmento", (), {"inicio": ticket.creado_at, "fin": ticket.cerrado_at, "cuenta_sla": True})()]
        for segmento in segmentos:
            inicio = max(segmento.inicio, inicio_rango)
            fin = min(segmento.fin or fin_rango, fin_rango)
            if inicio > fin:
                continue
            while inicio.date() <= fin.date():
                dia = inicio.date()
                fin_dia = timezone.make_aware(datetime.combine(dia, time.max))
                marca = "ATENCION" if segmento.cuenta_sla else "PAUSA"
                # Un ticket vencido no convierte en rojo sus días anteriores.
                # El día se marca excedido solamente si el segmento efectivo
                # tuvo tiempo posterior al instante real del vencimiento.
                if (
                    segmento.cuenta_sla
                    and ticket.sla_excedido_at
                    and min(fin, fin_dia) >= ticket.sla_excedido_at
                ):
                    marca = "EXCEDIDO"
                anterior = grupo["marcas"].get(dia)
                if anterior is None or PRIORIDAD_SLA_REPORTE[marca] > PRIORIDAD_SLA_REPORTE[anterior]:
                    grupo["marcas"][dia] = marca
                inicio = timezone.make_aware(datetime.combine(dia + timedelta(days=1), time.min))
    return dias, list(agrupadas.values())


@administrador_required
def reportería_incidencias(request):
    form = ReporteIncidenciasForm(request.GET or None)
    filas = _filas_reporte_incidencias(form.cleaned_data) if form.is_valid() else []
    tickets_unicos = {fila["ticket"].pk: fila["ticket"] for fila in filas}
    total_tickets = len(tickets_unicos)
    total_excedidos = sum(1 for fila in filas if fila["excedido"])
    total_cerrados = sum(1 for ticket in tickets_unicos.values() if ticket.cerrado_at)
    resumen_reporte = {
        "tickets": total_tickets,
        "excedidos": total_excedidos,
        "en_objetivo": max(0, total_tickets - total_excedidos),
        "cerrados": total_cerrados,
        "cumplimiento": round(((total_tickets - total_excedidos) / total_tickets) * 100) if total_tickets else 0,
        "efectivo": sum(fila["efectivo"] for fila in filas),
        "pausa": sum(fila["pausa"] for fila in filas),
    }
    if request.GET.get("exportar") == "xlsx" and form.is_valid():
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Incidencias SLA"
        dias, tiendas = _matriz_sla_reporte(filas, form.cleaned_data)
        primera_columna_dias, ultima_columna_dias = 4, 3 + len(dias)
        fin_titulo = get_column_letter(ultima_columna_dias + 7)
        hoja.merge_cells(f"D1:{fin_titulo}1")
        hoja["D1"] = "Incidencias operativas con ticket · Seguimiento SLA"
        hoja["D1"].fill = PatternFill("solid", fgColor="1E3A5F")
        hoja["D1"].font = Font(color="FFFFFF", bold=True, size=14)
        hoja["D1"].alignment = Alignment(horizontal="center", vertical="center")
        leyenda = [("EXCEDIDO", "SLA excedido"), ("ATENCION", "Atención efectiva (cuenta SLA)"), ("PAUSA", "Pausa por tercero (no cuenta SLA)")]
        for renglon, (marca, texto) in enumerate(leyenda, start=2):
            hoja.cell(renglon, 1, "X")
            hoja.cell(renglon, 1).fill = PatternFill("solid", fgColor=COLORES_SLA_REPORTE[marca])
            hoja.cell(renglon, 1).font = Font(color="FFFFFF", bold=True)
            hoja.merge_cells(start_row=renglon, start_column=2, end_row=renglon, end_column=3)
            hoja.cell(renglon, 2, texto)
        dias_semana = ("lun.", "mar.", "mié.", "jue.", "vie.", "sáb.", "dom.")
        for indice, dia in enumerate(dias, start=primera_columna_dias):
            hoja.cell(6, indice, dias_semana[dia.weekday()])
            hoja.cell(7, indice, dia.day)
        encabezados = ["Distrito", "Coordinador", "Tienda"]
        adicionales = ["Folios asociados", "Tipo", "Comentario", "SLA efectivo", "Pausa", "Límite SLA", "Estado SLA"]
        for indice, valor in enumerate(encabezados, start=1):
            hoja.cell(7, indice, valor)
        for indice, valor in enumerate(adicionales, start=ultima_columna_dias + 1):
            hoja.cell(7, indice, valor)
        for celda in hoja[6] + hoja[7]:
            celda.fill = PatternFill("solid", fgColor="315EE4")
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for tienda in tiendas:
            renglon = hoja.max_row + 1
            hoja.append([
                tienda["distrital"], tienda["coordinador"], tienda["tienda"],
                *["X" if dia in tienda["marcas"] else "" for dia in dias],
                " · ".join(dict.fromkeys(tienda["folios"])),
                " · ".join(dict.fromkeys(tienda["tipos"])),
                " · ".join(dict.fromkeys(tienda["comentarios"])) or "—",
                _formato_segundos(tienda["efectivo"]), _formato_segundos(tienda["pausa"]),
                " · ".join(f"{limite} min" for limite in sorted(set(tienda["limites"]))) if tienda["limites"] else "Sin definir",
                "SLA excedido" if tienda["excedido"] else "En objetivo",
            ])
            for indice, dia in enumerate(dias, start=primera_columna_dias):
                marca = tienda["marcas"].get(dia)
                if marca:
                    celda = hoja.cell(renglon, indice)
                    celda.fill = PatternFill("solid", fgColor=COLORES_SLA_REPORTE[marca])
                    celda.font = Font(color="FFFFFF", bold=True)
                    celda.alignment = Alignment(horizontal="center")
            estado = hoja.cell(renglon, ultima_columna_dias + len(adicionales))
            if tienda["excedido"]:
                estado.fill = PatternFill("solid", fgColor=COLORES_SLA_REPORTE["EXCEDIDO"])
                estado.font = Font(color="FFFFFF", bold=True)
        hoja.row_dimensions[1].height = 28
        hoja.freeze_panes = "D8"
        ultima_columna = ultima_columna_dias + len(adicionales)
        hoja.auto_filter.ref = f"A7:{get_column_letter(ultima_columna)}{hoja.max_row}"
        borde_suave = Border(bottom=Side(style="hair", color="D8E0EA"))
        for renglon in range(8, hoja.max_row + 1):
            if renglon % 2 == 0:
                for celda in hoja[renglon]:
                    if celda.fill.fill_type is None:
                        celda.fill = PatternFill("solid", fgColor="F5F8FC")
            for celda in hoja[renglon]:
                celda.border = borde_suave
                celda.alignment = Alignment(vertical="top", wrap_text=True)
        for columna in range(1, 4):
            hoja.column_dimensions[get_column_letter(columna)].width = (13, 22, 19)[columna - 1]
        for columna in range(primera_columna_dias, ultima_columna_dias + 1):
            hoja.column_dimensions[get_column_letter(columna)].width = 4.5
        for indice, ancho in enumerate((28, 30, 48, 18, 18, 14, 16), start=ultima_columna_dias + 1):
            hoja.column_dimensions[get_column_letter(indice)].width = ancho

        # Resumen entregable para medir desempeño sin mezclar pausas externas
        # con el trabajo efectivo del encargado.
        metricas = {}
        for fila in filas:
            ticket = fila["ticket"]
            nombre = _nombre_reporte_usuario(ticket.responsable)
            dato = metricas.setdefault(nombre, {"tickets": set(), "efectivo": 0, "pausa": 0, "cerrados": 0, "excedidos": 0})
            dato["tickets"].add(ticket.pk)
            dato["efectivo"] += fila["efectivo"]
            dato["pausa"] += fila["pausa"]
            dato["cerrados"] += int(ticket.cerrado_at is not None)
            dato["excedidos"] += int(fila["excedido"])
        hoja_metricas = libro.create_sheet("Métricas por encargado")
        columnas_metricas = ["Encargado", "Tickets atendidos", "Tiempo efectivo", "Pausa por tercero", "Cierres efectivos", "SLA excedidos", "% cumplimiento SLA"]
        hoja_metricas.append(columnas_metricas)
        for celda in hoja_metricas[1]:
            celda.fill = PatternFill("solid", fgColor="1E3A5F")
            celda.font = Font(color="FFFFFF", bold=True)
        for nombre, dato in sorted(metricas.items()):
            total = len(dato["tickets"])
            cumplimiento = (total - dato["excedidos"]) / total if total else 0
            hoja_metricas.append([nombre, total, _formato_segundos(dato["efectivo"]), _formato_segundos(dato["pausa"]), dato["cerrados"], dato["excedidos"], cumplimiento])
            hoja_metricas.cell(hoja_metricas.max_row, 7).number_format = "0.0%"
        hoja_metricas.freeze_panes = "A2"
        hoja_metricas.auto_filter.ref = hoja_metricas.dimensions
        for columna, ancho in enumerate((28, 18, 22, 22, 18, 16, 20), start=1):
            hoja_metricas.column_dimensions[get_column_letter(columna)].width = ancho

        hoja_resumen = libro.create_sheet("Resumen ejecutivo")
        hoja_resumen.sheet_view.showGridLines = False
        hoja_resumen.merge_cells("A1:D1")
        hoja_resumen["A1"] = "Resumen ejecutivo de SLA"
        hoja_resumen["A1"].fill = PatternFill("solid", fgColor="1E3A5F")
        hoja_resumen["A1"].font = Font(color="FFFFFF", bold=True, size=16)
        hoja_resumen["A1"].alignment = Alignment(horizontal="center", vertical="center")
        hoja_resumen.row_dimensions[1].height = 32
        periodo = f"{form.cleaned_data['fecha_inicio']:%d/%m/%Y} al {form.cleaned_data['fecha_fin']:%d/%m/%Y}"
        hoja_resumen.append(["Periodo", periodo])
        hoja_resumen.append([])
        indicadores = (
            ("Tickets en el periodo", resumen_reporte["tickets"]),
            ("En objetivo", resumen_reporte["en_objetivo"]),
            ("SLA excedidos", resumen_reporte["excedidos"]),
            ("Tickets cerrados", resumen_reporte["cerrados"]),
            ("Cumplimiento SLA", resumen_reporte["cumplimiento"] / 100),
            ("Tiempo efectivo", _formato_segundos(resumen_reporte["efectivo"])),
            ("Pausa por terceros", _formato_segundos(resumen_reporte["pausa"])),
        )
        for etiqueta, valor in indicadores:
            hoja_resumen.append([etiqueta, valor])
        hoja_resumen["B8"].number_format = "0%"
        for renglon in range(4, 11):
            hoja_resumen.cell(renglon, 1).font = Font(bold=True, color="42526A")
            hoja_resumen.cell(renglon, 1).fill = PatternFill("solid", fgColor="EEF3F8")
            hoja_resumen.cell(renglon, 2).font = Font(bold=True, color="172033")
        hoja_resumen.column_dimensions["A"].width = 28
        hoja_resumen.column_dimensions["B"].width = 26
        contenido = BytesIO()
        libro.save(contenido)
        respuesta = HttpResponse(contenido.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        respuesta["Content-Disposition"] = (
            'attachment; filename="reporte_sla_'
            f'{form.cleaned_data["fecha_inicio"]:%Y%m%d}_'
            f'{form.cleaned_data["fecha_fin"]:%Y%m%d}.xlsx"'
        )
        return respuesta
    return render(request, "tickets/reporte_incidencias.html", {
        "form": form,
        "filas": filas,
        "total_filas": len(filas),
        "resumen": resumen_reporte,
    })


def _formato_segundos(segundos):
    horas, resto = divmod(max(0, segundos) // 60, 60)
    dias, horas = divmod(horas, 24)
    return f"{dias} d {horas} h {resto} min"


@administrador_required
def estructura_operativa(request):
    """Centro de acceso; cada catálogo se administra en su propio panel."""
    return render(request, "tickets/estructura.html", {
        "resumen": {
            "empresas": Empresa.objects.count(),
            "zonas": Zona.objects.filter(activa=True).count(),
            "tiendas": Tienda.objects.count(),
            "equipos": GrupoTrabajo.objects.filter(activo=True).count(),
            "turnos": ProgramacionTrabajo.objects.count(),
            "sla": ConfiguracionSLA.objects.filter(activo=True).count(),
        },
    })


@administrador_required
def inventario(request):
    """Panel reservado para el futuro módulo de inventario."""
    return render(request, "tickets/inventario.html")


def _aplicar_busqueda(queryset, busqueda, campos):
    if not busqueda:
        return queryset
    condicion = Q()
    for campo in campos:
        condicion |= Q(**{f"{campo}__icontains": busqueda})
    return queryset.filter(condicion).distinct()


def _nombre_usuario(usuario):
    if not usuario:
        return "—"
    return usuario.get_full_name().strip() or usuario.username


@administrador_required
def panel_configuracion(request, seccion):
    """Compatibilidad con enlaces anteriores: los catálogos ahora están agrupados."""
    destinos = {"empresas": "territorio", "zonas": "territorio", "tiendas": "territorio",
                "calendario": "calendario", "sla": "incidencias", "coberturas": "estructura"}
    if seccion not in destinos:
        raise Http404("Configuración no encontrada.")
    return redirect("tickets:" + destinos[seccion])


@administrador_required
def panel_equipos(request):
    return redirect("tickets:calendario")


@administrador_required
def centro_importaciones(request):
    recientes = (
        ImportacionTickets.objects
        .select_related("empresa", "usuario")
        .order_by("-iniciado_at")[:12]
    )
    return render(request, "tickets/centro_importaciones.html", {
        "recientes": recientes,
    })


@administrador_required
def directorio_activos(request):
    consulta = request.GET.get("q", "").strip()
    activos = AsignacionPersonal.objects.filter(
        funcion=AsignacionPersonal.Funcion.PERSONAL,
        activa=True,
    ).select_related("persona", "tienda", "empresa").order_by("tienda__nombre", "persona__nombre")
    if consulta:
        activos = activos.filter(Q(persona__nombre__icontains=consulta) | Q(persona__id_personal__icontains=consulta) | Q(tienda__nombre__icontains=consulta) | Q(tienda__clave__icontains=consulta))
    pagina = Paginator(activos, 50).get_page(request.GET.get("page"))
    return render(request, "tickets/directorio_activos.html", {"pagina": pagina, "consulta": consulta})


@administrador_required
def importar_catalogos_view(request):
    previsualizacion = request.session.get("catalogo_preview")
    tipo_solicitado = request.GET.get("tipo", "")
    initial = (
        {"tipo": tipo_solicitado}
        if tipo_solicitado in dict(ImportacionCatalogoForm.TIPO_CHOICES)
        else None
    )
    form = ImportacionCatalogoForm(
        request.POST or None,
        request.FILES or None,
        initial=initial,
    )
    if request.method == "POST" and request.POST.get("accion") == "analizar" and form.is_valid():
        try:
            empresa = form.cleaned_data.get("empresa")
            previsualizacion = analizar_catalogo(
                form.cleaned_data["archivo"],
                form.cleaned_data["tipo"],
                empresa=empresa,
            )
            previsualizacion["empresa_id"] = empresa.pk if empresa else None
            request.session["catalogo_preview"] = previsualizacion
        except Exception as error:
            messages.error(request, f"No se pudo analizar el archivo: {error}")
            previsualizacion = None
    elif request.method == "POST" and request.POST.get("accion") == "confirmar" and previsualizacion:
        empresa = Empresa.objects.filter(pk=previsualizacion.get("empresa_id")).first()
        try:
            resultado = importar_catalogo(previsualizacion, empresa)
            enviados = sum(enviar_invitacion_usuario(usuario, request) for usuario in resultado.pop("usuarios_nuevos"))
            request.session.pop("catalogo_preview", None)
            messages.success(request, f"Importación completada: {resultado['nuevos']} nuevos, {resultado['actualizados']} actualizados y {resultado['sin_cambios']} sin cambios. Invitaciones enviadas: {enviados}.")
            return redirect("tickets:estructura")
        except Exception as error:
            messages.error(request, f"No se pudo confirmar la importación: {error}")
    return render(request, "tickets/importar_catalogos.html", {"form": form, "resultado": previsualizacion})


@administrador_required
def descargar_plantilla_tiendas(request):
    encabezados = [
        "ID de tienda", "Key", "Tienda", "Nombre de la tienda", "Dirección",
        "Estado", "Agencia", "Socio", "Distrital", "Gerente distrital",
        "Coordinador", "Tipo de coordinador", "Encargado de soporte",
        "Partner", "Líder actual", "Retiro", "Prioridad", "Piloto",
        "Latitud", "Longitud",
    ]
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tiendas"
    hoja.append(encabezados)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}1"
    relleno = PatternFill("solid", fgColor="315EE4")
    for indice, celda in enumerate(hoja[1], start=1):
        celda.fill = relleno
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(indice)].width = max(16, len(celda.value) + 4)
    hoja.row_dimensions[1].height = 24

    instrucciones = libro.create_sheet("Instrucciones")
    instrucciones.append(["Campo", "Uso"])
    instrucciones.append(["ID de tienda", "Obligatorio. Identificador externo único por empresa."])
    instrucciones.append(["Tienda", "Obligatorio. Número o referencia operativa."])
    instrucciones.append(["Nombre de la tienda", "Obligatorio. Nombre visible de la sucursal."])
    instrucciones.append(["Estado", "Obligatorio. Estado de México."])
    instrucciones.append(["Distrital", "Obligatorio. Distrito operativo."])
    instrucciones.append(["Retiro / Piloto", "Usar Sí o No. Si se dejan vacíos se interpretan como No."])
    instrucciones.append(["Latitud / Longitud", "Opcionales. Deben capturarse juntas; longitud es negativa en México."])
    instrucciones.append(["Código interno", "No se captura: el sistema genera TDA-###### automáticamente."])
    instrucciones.column_dimensions["A"].width = 24
    instrucciones.column_dimensions["B"].width = 75
    for celda in instrucciones[1]:
        celda.fill = relleno
        celda.font = Font(color="FFFFFF", bold=True)

    contenido = BytesIO()
    libro.save(contenido)
    respuesta = HttpResponse(
        contenido.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = 'attachment; filename="plantilla_tiendas_carobra.xlsx"'
    return respuesta


@administrador_required
def descargar_plantilla_coordenadas(request):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Coordenadas"
    hoja.append(["ID de tienda", "Latitud", "Longitud"])
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = "A1:C1"
    relleno = PatternFill("solid", fgColor="315EE4")
    anchos = (24, 18, 18)
    for indice, celda in enumerate(hoja[1], start=1):
        celda.fill = relleno
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(indice)].width = anchos[indice - 1]
    hoja.row_dimensions[1].height = 24
    contenido = BytesIO()
    libro.save(contenido)
    respuesta = HttpResponse(
        contenido.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = 'attachment; filename="plantilla_coordenadas_tiendas.xlsx"'
    return respuesta


@administrador_required
def actualizar_coordenadas_tienda(request):
    if request.method != "POST":
        return redirect("tickets:mapa_tiendas")
    form = CoordenadasTiendaForm(request.POST)
    if form.is_valid():
        tienda = form.cleaned_data["tienda"]
        tienda.latitud = form.cleaned_data["latitud"]
        tienda.longitud = form.cleaned_data["longitud"]
        tienda.estado_geocodificacion = Tienda.EstadoGeocodificacion.LOCALIZADA
        tienda.detalle_geocodificacion = "Coordenadas actualizadas manualmente."
        tienda.geocodificada_at = timezone.now()
        tienda.save(update_fields=[
            "latitud", "longitud", "estado_geocodificacion",
            "detalle_geocodificacion", "geocodificada_at", "actualizado_at",
        ])
        messages.success(request, f"Ubicación de {tienda.nombre} actualizada correctamente.")
    else:
        error = next(iter(form.errors.values()))[0]
        messages.error(request, f"No se actualizó la ubicación: {error}")
    return redirect("tickets:mapa_tiendas")


@administrador_required
def mapa_tiendas(request):
    tiendas = Tienda.objects.select_related("empresa", "zona").order_by("nombre")
    busqueda = request.GET.get("q", "").strip()
    filtros = {
        "empresa": request.GET.get("empresa", "").strip(),
        "estado": request.GET.get("estado", "").strip(),
        "ciudad": request.GET.get("ciudad", "").strip(),
        "distrital": request.GET.get("distrital", "").strip(),
        "agencia": request.GET.get("agencia", "").strip(),
        "socio": request.GET.get("socio", "").strip(),
        "coordinador": request.GET.get("coordinador", "").strip(),
        "tipo_coordinador": request.GET.get("tipo_coordinador", "").strip(),
        "gerente": request.GET.get("gerente", "").strip(),
        "encargado": request.GET.get("encargado", "").strip(),
        "partner": request.GET.get("partner", "").strip(),
        "prioridad": request.GET.get("prioridad", "").strip(),
        "piloto": request.GET.get("piloto", "").strip(),
        "retiro": request.GET.get("retiro", "").strip(),
    }
    if busqueda:
        tiendas = tiendas.filter(
            Q(codigo__icontains=busqueda)
            | Q(id_externo__icontains=busqueda)
            | Q(clave__icontains=busqueda)
            | Q(numero__icontains=busqueda)
            | Q(nombre__icontains=busqueda)
            | Q(ciudad__icontains=busqueda)
            | Q(direccion__icontains=busqueda)
        )
    campos = {
        "empresa": "empresa_id",
        "estado": "estado",
        "ciudad": "ciudad",
        "distrital": "zona__nombre",
        "agencia": "agencia",
        "socio": "socio",
        "coordinador": "coordinador",
        "tipo_coordinador": "tipo_coordinador",
        "gerente": "gerente_distrital",
        "encargado": "encargado_soporte_nombre",
        "partner": "partner_nombre",
        "prioridad": "prioridad_operativa",
    }
    for parametro, campo in campos.items():
        if parametro == "ciudad" and filtros[parametro] == "__SIN_CIUDAD__":
            tiendas = tiendas.filter(ciudad="")
        elif filtros[parametro]:
            tiendas = tiendas.filter(**{campo: filtros[parametro]})
    for parametro in ("piloto", "retiro"):
        if filtros[parametro] in {"si", "no"}:
            tiendas = tiendas.filter(**{parametro: filtros[parametro] == "si"})

    base = Tienda.objects.select_related("zona")
    opciones = {
        parametro: list(
            base.exclude(**{f"{campo}__exact": ""})
            .values_list(campo, flat=True)
            .distinct()
            .order_by(campo)
        )
        for parametro, campo in campos.items()
        if parametro != "empresa"
    }
    opciones["empresa"] = list(Empresa.objects.order_by("nombre").values("id", "nombre"))
    marcadores = [
        {
            "codigo": tienda.codigo,
            "id_externo": tienda.id_externo,
            "numero": tienda.numero,
            "nombre": tienda.nombre,
            "direccion": tienda.direccion,
            "estado": tienda.estado,
            "ciudad": tienda.ciudad,
            "distrital": tienda.zona.nombre,
            "agencia": tienda.agencia,
            "socio": tienda.socio,
            "cadena": tienda.tipo_tienda_codigo,
            "gerente": tienda.gerente_distrital,
            "coordinador": tienda.coordinador,
            "encargado": tienda.encargado_soporte_nombre,
            "partner": tienda.partner_nombre,
            "prioridad": tienda.prioridad_operativa,
            "empresa": tienda.empresa.nombre,
            "piloto": tienda.piloto,
            "retiro": tienda.retiro,
            "latitud": float(tienda.latitud),
            "longitud": float(tienda.longitud),
            "id": tienda.pk,
        }
        for tienda in tiendas
        if tienda.latitud is not None and tienda.longitud is not None
    ]
    pendientes_query = tiendas.filter(
        Q(latitud__isnull=True) | Q(longitud__isnull=True)
    )
    pendientes_total = pendientes_query.count()
    pendientes_tiendas = list(
        pendientes_query.values("id", "nombre", "direccion", "estado", "ciudad")[:12]
    )
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse({
            "marcadores": marcadores,
            "total": tiendas.count(),
            "localizadas": len(marcadores),
            "pendientes": pendientes_total,
            "pendientes_tiendas": pendientes_tiendas,
        })
    coordenadas_catalogo = list(
        Tienda.objects.order_by("empresa__nombre", "nombre").values(
            "id", "latitud", "longitud"
        )
    )
    for item in coordenadas_catalogo:
        item["latitud"] = float(item["latitud"]) if item["latitud"] is not None else None
        item["longitud"] = float(item["longitud"]) if item["longitud"] is not None else None
    return render(request, "tickets/mapa_tiendas.html", {
        "tiendas": tiendas,
        "marcadores": marcadores,
        "total": tiendas.count(),
        "localizadas": len(marcadores),
        "pendientes": pendientes_total,
        "pendientes_tiendas": pendientes_tiendas,
        "busqueda": busqueda,
        "filtros": filtros,
        "opciones": opciones,
        "coordenadas_form": CoordenadasTiendaForm(),
        "coordenadas_catalogo": coordenadas_catalogo,
        "tienda_seleccionada": request.GET.get("tienda", "").strip(),
    })


def _leer_rango(path, inicio, longitud, bloque=1024 * 1024):
    with path.open("rb") as archivo:
        archivo.seek(inicio)
        restante = longitud
        while restante:
            contenido = archivo.read(min(bloque, restante))
            if not contenido:
                break
            restante -= len(contenido)
            yield contenido


@administrador_required
def mapa_vectorial_mexico(request):
    """Sirve el PMTiles local con soporte de rangos para MapLibre."""
    path = Path(settings.MEDIA_ROOT) / "mapas" / "mexico-20260913.pmtiles"
    if not path.exists():
        raise Http404("El mapa vectorial de México todavía no está disponible.")
    total = path.stat().st_size
    rango = request.headers.get("range", "")
    coincidencia = re.fullmatch(r"bytes=(\d+)-(\d*)", rango)
    inicio, fin, estado = 0, total - 1, 200
    if coincidencia:
        inicio = int(coincidencia.group(1))
        fin = int(coincidencia.group(2)) if coincidencia.group(2) else total - 1
        if inicio >= total or fin < inicio:
            respuesta = HttpResponse(status=416)
            respuesta["Content-Range"] = f"bytes */{total}"
            return respuesta
        fin = min(fin, total - 1)
        estado = 206
    longitud = fin - inicio + 1
    if request.method == "HEAD":
        respuesta = HttpResponse(status=estado, content_type="application/vnd.pmtiles")
    else:
        respuesta = StreamingHttpResponse(
            _leer_rango(path, inicio, longitud),
            status=estado,
            content_type="application/vnd.pmtiles",
        )
    respuesta["Accept-Ranges"] = "bytes"
    respuesta["Content-Length"] = str(longitud)
    respuesta["Cache-Control"] = "private, max-age=86400"
    respuesta["ETag"] = f'"{path.stat().st_mtime_ns:x}-{total:x}"'
    if estado == 206:
        respuesta["Content-Range"] = f"bytes {inicio}-{fin}/{total}"
    return respuesta


@administrador_required
def geometria_mexico(request, archivo):
    permitidos = {
        "estados": "mexico-estados-simplificado.geojson",
        "mascara": "mexico-mask.geojson",
    }
    nombre = permitidos.get(archivo)
    if not nombre:
        raise Http404
    path = Path(settings.MEDIA_ROOT) / "mapas" / nombre
    if not path.exists():
        raise Http404("No se encontró la geometría de México.")
    comprimido = Path(f"{path}.gz")
    if "gzip" in request.headers.get("accept-encoding", "") and comprimido.exists():
        respuesta = FileResponse(comprimido.open("rb"), content_type="application/geo+json")
        respuesta["Content-Encoding"] = "gzip"
        respuesta["Vary"] = "Accept-Encoding"
        respuesta["Cache-Control"] = "private, max-age=86400"
        return respuesta
    respuesta = FileResponse(path.open("rb"), content_type="application/geo+json")
    respuesta["Cache-Control"] = "private, max-age=86400"
    return respuesta


@administrador_required
def importar_bradescard_view(request):
    """
    Pantalla para cargar y analizar un Excel Bradescard.

    En esta vista todavía NO se modifican tickets.
    Solo se guarda temporalmente el archivo y
    se ejecuta la previsualización / dry run.
    """

    resultado = None
    error = None

    if request.method == "POST":
        form = ImportacionBradescardForm(
            request.POST,
            request.FILES,
        )

        if form.is_valid():
            archivo = form.cleaned_data["archivo"]
            empresa = form.cleaned_data["empresa"]

            # =================================================
            # CARPETA TEMPORAL
            # =================================================

            carpeta_tmp = (
                Path(settings.MEDIA_ROOT)
                / "importaciones_tmp"
            )

            carpeta_tmp.mkdir(
                parents=True,
                exist_ok=True,
            )

            # =================================================
            # ELIMINAR PREVISUALIZACIÓN ANTERIOR
            # =================================================

            anterior = request.session.get(
                "bradescard_preview_path"
            )

            if anterior:
                ruta_anterior = Path(anterior)

                try:
                    if (
                        ruta_anterior.exists()
                        and carpeta_tmp in ruta_anterior.parents
                    ):
                        ruta_anterior.unlink()

                except OSError:
                    pass

            # =================================================
            # GUARDAR ARCHIVO TEMPORAL
            # =================================================

            nombre_seguro = (
                f"{uuid4().hex}.xlsx"
            )

            ruta_temporal = (
                carpeta_tmp
                / nombre_seguro
            )

            try:
                with open(
                    ruta_temporal,
                    "wb",
                ) as destino:

                    for bloque in archivo.chunks():
                        destino.write(bloque)

                # =============================================
                # ANALIZAR SIN GUARDAR TICKETS
                # =============================================

                resultado = (
                    analizar_importacion_bradescard(
                        ruta_temporal,
                        empresa=empresa,
                    )
                )

                resultado["nombre_original"] = (
                    archivo.name
                )
                resultado["empresa"] = empresa.nombre

                # =============================================
                # GUARDAR PREVISUALIZACIÓN EN SESIÓN
                # =============================================

                request.session[
                    "bradescard_preview_path"
                ] = str(ruta_temporal)

                request.session[
                    "bradescard_preview_name"
                ] = archivo.name
                request.session["bradescard_preview_empresa_id"] = empresa.pk
                request.session["bradescard_preview_categorias"] = resultado["categorias_seleccionadas"]
                request.session["bradescard_categorias_disponibles"] = resultado["categorias_disponibles"]

            except Exception as exc:
                error = str(exc)

                try:
                    if ruta_temporal.exists():
                        ruta_temporal.unlink()

                except OSError:
                    pass

                request.session.pop(
                    "bradescard_preview_path",
                    None,
                )

                request.session.pop(
                    "bradescard_preview_name",
                    None,
                )

    else:
        form = ImportacionBradescardForm()

    return render(
        request,
        "tickets/importar_bradescard.html",
        {
            "form": form,
            "resultado": resultado,
            "error": error,
        },
    )


@administrador_required
def confirmar_importacion_bradescard(request):
    """
    Ejecuta la importación REAL del archivo
    que previamente fue analizado.

    Solo acepta POST.
    """

    # =====================================================
    # SOLO POST
    # =====================================================

    if request.method != "POST":
        return redirect(
            "tickets:importar_bradescard"
        )

    # =====================================================
    # RECUPERAR ARCHIVO ANALIZADO
    # =====================================================

    ruta = request.session.get(
        "bradescard_preview_path"
    )

    nombre_original = request.session.get(
        "bradescard_preview_name"
    )
    empresa = get_object_or_404(
        Empresa,
        pk=request.session.get("bradescard_preview_empresa_id"),
    )

    if not ruta:
        messages.error(
            request,
            (
                "No existe un archivo analizado "
                "pendiente de importar."
            ),
        )

        return redirect(
            "tickets:importar_bradescard"
        )

    ruta_archivo = Path(ruta)

    # =====================================================
    # VALIDAR QUE EL ARCHIVO SIGA EXISTIENDO
    # =====================================================

    if not ruta_archivo.exists():

        messages.error(
            request,
            (
                "El archivo temporal ya no existe. "
                "Vuelve a analizar la plantilla."
            ),
        )

        request.session.pop(
            "bradescard_preview_path",
            None,
        )

        request.session.pop(
            "bradescard_preview_name",
            None,
        )

        return redirect(
            "tickets:importar_bradescard"
        )

    # =====================================================
    # EJECUTAR IMPORTACIÓN
    # =====================================================

    categorias = request.POST.getlist("categorias")
    disponibles = request.session.get("bradescard_categorias_disponibles", [])
    if any(c not in disponibles for c in categorias):
        messages.error(request, "La selección de categorías no corresponde al archivo analizado.")
        return redirect("tickets:importar_bradescard")
    if sorted(categorias) != sorted(request.session.get("bradescard_preview_categorias", [])) or request.POST.get("accion") == "filtrar":
        resultado = analizar_importacion_bradescard(ruta_archivo, empresa=empresa, categorias=categorias)
        resultado["nombre_original"] = nombre_original
        request.session["bradescard_preview_categorias"] = categorias
        return render(request, "tickets/importar_bradescard.html", {"form": ImportacionBradescardForm(initial={"empresa": empresa}), "resultado": resultado})
    if not categorias:
        messages.error(request, "Selecciona al menos una categoría.")
        return redirect("tickets:importar_bradescard")
    try:
        resultado = importar_bradescard(
            ruta_archivo,
            usuario=request.user,
            empresa=empresa,
            categorias=categorias,
        )

    except Exception as exc:

        messages.error(
            request,
            f"Error durante la importación: {exc}",
        )

        return redirect(
            "tickets:importar_bradescard"
        )

    # =====================================================
    # IMPORTACIÓN EXITOSA
    # =====================================================

    if resultado.get("ok"):

        messages.success(
            request,
            (
                f"Importación completada correctamente. "
                f"Nuevos: {resultado['nuevos']} · "
                f"Actualizados: {resultado['actualizados']} · "
                f"Sin cambios: {resultado['sin_cambios']}."
            ),
        )

        # ---------------------------------------------
        # Eliminar archivo temporal
        # ---------------------------------------------

        try:
            ruta_archivo.unlink()

        except OSError:
            pass

        # ---------------------------------------------
        # Limpiar sesión
        # ---------------------------------------------

        request.session.pop(
            "bradescard_preview_path",
            None,
        )

        request.session.pop(
            "bradescard_preview_name",
            None,
        )

        return redirect(
            "tickets:importar_bradescard"
        )

    # =====================================================
    # ARCHIVO DUPLICADO
    # =====================================================

    if resultado.get("duplicado"):

        messages.warning(
            request,
            (
                f"El archivo "
                f"{nombre_original or 'seleccionado'} "
                f"ya fue importado anteriormente."
            ),
        )

        return redirect(
            "tickets:importar_bradescard"
        )

    # =====================================================
    # ERROR DEVUELTO POR EL IMPORTADOR
    # =====================================================

    messages.error(
        request,
        resultado.get(
            "error",
            "No se pudo completar la importación.",
        ),
    )

    return redirect(
        "tickets:importar_bradescard"
    )

@login_required
def crear_ticket_carobra(request):
    form = TicketCarobraForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        tienda = form.cleaned_data["tienda_registrada"]
        responsable = responsable_automatico(tienda, form.configuracion_sla())
        partner = partner_automatico(tienda)

        with transaction.atomic():
            ticket = form.save(
                creado_por=request.user,
                responsable=responsable,
                partner=partner,
            )

            if responsable:
                ticket.asignado_at = timezone.now()
                ticket.save(update_fields=["asignado_at", "actualizado_at"])

            HistorialTicket.objects.create(
                ticket=ticket,
                usuario=request.user,
                evento=HistorialTicket.Evento.CREADO,
                origen=HistorialTicket.Origen.USUARIO,
                descripcion="Ticket CAROBRA creado manualmente.",
                valor_nuevo={
                    "origen": OrigenTicket.CAROBRA,
                    "estado": ticket.estado_interno,
                    "tienda": ticket.tienda,
                    "prioridad": ticket.prioridad,
                    "sla_limite_minutos": ticket.sla_limite_minutos,
                },
            )

            if responsable:
                HistorialTicket.objects.create(
                    ticket=ticket,
                    usuario=None,
                    evento=HistorialTicket.Evento.ASIGNADO,
                    origen=HistorialTicket.Origen.SISTEMA,
                    descripcion="Responsable asignado automáticamente por tienda.",
                    valor_nuevo={"responsable": responsable.username},
                )

        if ticket.configuracion_sla_id:
            messages.success(
                request,
                "Ticket CAROBRA creado con prioridad y SLA automáticos.",
            )
        else:
            messages.warning(
                request,
                "Ticket creado. La incidencia aún no tiene un SLA configurado.",
            )

        return redirect("tickets:detalle", folio=ticket.folio)

    configuraciones = (
        ConfiguracionSLA.objects
        .filter(activo=True)
        .order_by("categoria", "incidencia_general")
    )
    return render(
        request,
        "tickets/crear_ticket.html",
        {"form": form, "configuraciones_sla": configuraciones},
    )


@login_required
def lista_tickets(request):
    """
    Listado principal de tickets.

    - 50 registros por página.
    - Búsqueda general.
    - Filtros por estado, categoría y tipo.
    - No carga los 4,625 registros de golpe.
    """

    # =====================================================
    # PARÁMETROS
    # =====================================================

    busqueda = request.GET.get(
        "q",
        "",
    ).strip()

    estado = request.GET.get(
        "estado",
        "",
    ).strip()

    categoria = request.GET.get(
        "categoria",
        "",
    ).strip()

    tipo = request.GET.get(
        "tipo",
        "",
    ).strip()

    origen = request.GET.get("origen", "").strip()
    prioridad = request.GET.get("prioridad", "").strip()
    vista = request.GET.get("vista", "").strip()
    if vista not in {"encargado", "partner", "todos"}:
        vista = "todos" if es_administrador(request.user) else "encargado"
    fecha_apertura_desde = request.GET.get("fecha_apertura_desde", "").strip()
    fecha_apertura_hasta = request.GET.get("fecha_apertura_hasta", "").strip()
    fecha_cierre_desde = request.GET.get("fecha_cierre_desde", "").strip()
    fecha_cierre_hasta = request.GET.get("fecha_cierre_hasta", "").strip()
    orden = request.GET.get("orden", "").strip()
    direccion = request.GET.get("direccion", "").strip()

    campos_ordenables = {
        "ticket": "ticket_bradescard",
        "tienda": "tienda",
        "categoria": "categoria",
        "tipo": "tipo_externo",
        "estado": "estado_interno",
        "prioridad": "prioridad",
        "responsable": "responsable__username",
        "sla": "sla_vencido_actual",
    }

    if orden not in campos_ordenables:
        orden = "reciente"

    if direccion not in {"asc", "desc"}:
        direccion = "desc"


    # =====================================================
    # QUERY BASE
    # =====================================================

    tickets_base = con_sla_actual(tickets_visibles_para(request.user))
    if request.GET.get("requiere_admin") == "1":
        tickets_base = requieren_administracion(tickets_base)
    tickets = tickets_base
    if request.GET.get("empresa", "").isdigit():
        tickets = tickets.filter(empresa_id=int(request.GET["empresa"]))
    if request.GET.get("tienda_id", "").isdigit():
        tickets = tickets.filter(tienda_registrada_id=int(request.GET["tienda_id"]))
    if request.GET.get("abiertos") == "1":
        tickets = tickets.exclude(estado_interno=EstadoTicket.CERRADO)
    if vista == "encargado":
        tickets = tickets.filter(responsable=request.user)
    elif vista == "partner":
        tickets = tickets.filter(partner=request.user)
    if fecha_apertura_desde:
        tickets = tickets.filter(fecha_apertura_externa__date__gte=fecha_apertura_desde)
    if fecha_apertura_hasta:
        tickets = tickets.filter(fecha_apertura_externa__date__lte=fecha_apertura_hasta)
    if fecha_cierre_desde:
        tickets = tickets.filter(cerrado_at__date__gte=fecha_cierre_desde)
    if fecha_cierre_hasta:
        tickets = tickets.filter(cerrado_at__date__lte=fecha_cierre_hasta)


    # =====================================================
    # BÚSQUEDA
    # =====================================================

    if busqueda:
        tickets = tickets.filter(
            Q(folio__icontains=busqueda)
            | Q(ticket_bradescard__icontains=busqueda)
            | Q(tienda__icontains=busqueda)
            | Q(categoria__icontains=busqueda)
            | Q(incidencia_general__icontains=busqueda)
            | Q(incidencia_especifica__icontains=busqueda)
            | Q(usuario_externo__icontains=busqueda)
            | Q(asignado_externo__icontains=busqueda)
        )


    # =====================================================
    # FILTROS
    # =====================================================

    if estado:
        tickets = tickets.filter(
            estado_interno=estado
        )

    if categoria:
        tickets = tickets.filter(
            categoria=categoria
        )

    if tipo:
        tickets = tickets.filter(
            tipo_externo=tipo
        )

    if prioridad in PrioridadTicket.values:
        tickets = tickets.filter(prioridad=prioridad)

    if origen in OrigenTicket.values:
        tickets = tickets.filter(origen=origen)


    # =====================================================
    # ORDEN
    # =====================================================

    campo_orden = campos_ordenables.get(orden, "id")
    prefijo_direccion = "-" if direccion == "desc" else ""
    tickets = tickets.annotate(
        cerrado_al_final=Case(
            When(estado_interno=EstadoTicket.CERRADO, then=Value(1)),
            default=Value(0),
            output_field=IntegerField(),
        )
    ).order_by(
        "cerrado_al_final",
        f"{prefijo_direccion}{campo_orden}",
        "-id",
    )


    # =====================================================
    # PAGINACIÓN
    # =====================================================

    paginator = Paginator(
        tickets,
        50,
    )

    pagina = paginator.get_page(
        request.GET.get("page")
    )

    parametros_paginacion = request.GET.copy()
    parametros_paginacion.pop("page", None)

    parametros_orden = parametros_paginacion.copy()
    parametros_orden.pop("orden", None)
    parametros_orden.pop("direccion", None)

    contexto_resultados = {
        "pagina": pagina,
        "orden_actual": orden,
        "direccion_orden": direccion,
        "parametros_paginacion": parametros_paginacion.urlencode(),
        "parametros_orden": parametros_orden.urlencode(),
        "paginas_visibles": paginator.get_elided_page_range(pagina.number, on_each_side=2, on_ends=1),
    }

    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return render(
            request,
            "tickets/_lista_resultados.html",
            contexto_resultados,
        )


    # =====================================================
    # ESTADÍSTICAS GENERALES
    # =====================================================

    estadisticas = tickets_base.aggregate(
        total=Count("id"),

        nuevos=Count(
            "id",
            filter=Q(
                estado_interno=EstadoTicket.NUEVO
            ),
        ),

        asignados=Count(
            "id",
            filter=Q(
                estado_interno=EstadoTicket.ASIGNADO
            ),
        ),

        en_proceso=Count(
            "id",
            filter=Q(
                estado_interno=EstadoTicket.EN_PROCESO
            ),
        ),

        en_espera=Count(
            "id",
            filter=Q(
                estado_interno=EstadoTicket.EN_ESPERA
            ),
        ),

        resueltos=Count(
            "id",
            filter=Q(
                estado_interno=EstadoTicket.RESUELTO
            ),
        ),

        cerrados=Count(
            "id",
            filter=Q(
                estado_interno=EstadoTicket.CERRADO
            ),
        ),

        sla_excedidos=Count(
            "id",
            filter=Q(
                sla_vencido_actual=True
            ),
        ),

        carobra=Count(
            "id",
            filter=Q(origen=OrigenTicket.CAROBRA),
        ),

        bradescard=Count(
            "id",
            filter=Q(origen=OrigenTicket.BRADESCARD),
        ),
    )


    # =====================================================
    # CATÁLOGOS PARA FILTROS
    # =====================================================

    categorias = (
        tickets_base
        .exclude(categoria="")
        .values_list(
            "categoria",
            flat=True,
        )
        .distinct()
        .order_by("categoria")
    )

    tipos = (
        tickets_base
        .exclude(tipo_externo="")
        .values_list(
            "tipo_externo",
            flat=True,
        )
        .distinct()
        .order_by("tipo_externo")
    )


    # =====================================================
    # RENDER
    # =====================================================

    return render(
        request,
        "tickets/lista_tickets.html",
        {
            **contexto_resultados,
            "estadisticas": estadisticas,

            "categorias": categorias,
            "tipos": tipos,

            "busqueda": busqueda,
            "estado_actual": estado,
            "categoria_actual": categoria,
            "tipo_actual": tipo,
            "origen_actual": origen,
            "prioridad_actual": prioridad,
            "vista_actual": vista,
            "fecha_apertura_desde": fecha_apertura_desde,
            "fecha_apertura_hasta": fecha_apertura_hasta,
            "fecha_cierre_desde": fecha_cierre_desde,
            "fecha_cierre_hasta": fecha_cierre_hasta,

            "estados": EstadoTicket.choices,
            "origenes": OrigenTicket.choices,
            "prioridades": PrioridadTicket.choices,
        },
    )

@login_required
def detalle_ticket(request, folio):
    """
    Centro de gestión del ticket.

    Soporta:
    - Asignación rápida de responsable.
    - Gestión completa del ticket.
    - Historial por cada cambio.
    - Timestamps de asignación, resolución y cierre.
    - Reaperturas.
    - NUEVO + responsable -> ASIGNADO automáticamente.
    """

    puede_gestionar = es_administrador(request.user)

    ticket = get_object_or_404(
        tickets_visibles_para(request.user).select_related(
            "empresa",
            "creado_por",
            "responsable",
            "partner",
            "suplente",
            "tienda_registrada__zona",
            "tienda_registrada__encargado_soporte__usuario",
            "tienda_registrada__partner__usuario",
        ),
        folio=folio,
    )
    tienda_personal = ticket.tienda_registrada
    if tienda_personal is None and ticket.tienda:
        tienda_personal = Tienda.objects.filter(
            Q(clave__iexact=ticket.tienda)
            | Q(id_externo__iexact=ticket.tienda)
            | Q(numero__iexact=ticket.tienda)
            | Q(nombre__iexact=ticket.tienda)
        ).first()
    activos_tienda = (
        tienda_personal.asignaciones_personal.filter(
            funcion=AsignacionPersonal.Funcion.PERSONAL, activa=True
        ).select_related("persona") if tienda_personal else AsignacionPersonal.objects.none()
    )

    partner_usuario = ticket.partner

    es_responsable = ticket.responsable_id == request.user.id
    es_partner = bool(
        partner_usuario and partner_usuario.pk == request.user.id
    )
    es_creador = ticket.creado_por_id == request.user.id

    if puede_gestionar:
        rol_en_ticket = "Administrador"
    elif es_responsable:
        rol_en_ticket = "Responsable"
    elif es_partner:
        rol_en_ticket = "Partner"
    elif es_creador:
        rol_en_ticket = "Creador"
    else:
        rol_en_ticket = "Participante"

    # =====================================================
    # FORMULARIOS DE COMENTARIO / ARCHIVO
    # =====================================================

    comentario_form = ComentarioTicketForm(
        permitir_nota_interna=puede_gestionar
    )
    archivo_form = ArchivoTicketForm()
    cierre_form = CerrarTicketForm()
    reapertura_form = ReabrirTicketForm()

    # =====================================================
    # HELPERS
    # =====================================================

    def usuario_valor(usuario):
        if not usuario:
            return None

        return usuario.username

    def objeto_valor(objeto):
        if not objeto:
            return None

        return str(objeto)

    # =====================================================
    # POST
    # =====================================================

    if request.method == "POST":

        accion = request.POST.get(
            "accion",
            "responsable",
        )
        if not participa_en_ticket(request.user, ticket):
            raise PermissionDenied

        if accion not in {
            "comentario",
            "cierre",
            "reapertura",
            "gestion",
            "responsable",
        }:
            raise PermissionDenied

        if accion in {"gestion", "responsable"} and not puede_gestionar:
            if accion != "responsable" or not permiso_ticket(request.user, "puede_reasignar"):
                raise PermissionDenied

        if accion == "reapertura" and not puede_gestionar:
            raise PermissionDenied
        if accion == "cierre" and not permiso_ticket(request.user, "puede_cerrar"):
            raise PermissionDenied
        if accion == "comentario" and not permiso_ticket(request.user, "puede_seguimiento"):
            raise PermissionDenied

        # =================================================
        # COMENTARIO / SEGUIMIENTO
        # =================================================

        if accion == "comentario":

            comentario_form = ComentarioTicketForm(
                request.POST,
                permitir_nota_interna=puede_gestionar,
            )

            archivo_form = ArchivoTicketForm(
                request.POST,
                request.FILES,
            )

            if (
                comentario_form.is_valid()
                and archivo_form.is_valid()
            ):

                with transaction.atomic():

                    ticket = (
                        Ticket.objects
                        .select_for_update()
                        .get(pk=ticket.pk)
                    )

                    # -------------------------------------
                    # PRIMER COMENTARIO
                    # -------------------------------------

                    es_primer_comentario = (
                        ticket.primer_comentario_at is None
                    )

                    # -------------------------------------
                    # CREAR COMENTARIO
                    # -------------------------------------

                    comentario = comentario_form.save(
                        commit=False
                    )

                    comentario.ticket = ticket
                    comentario.usuario = request.user
                    comentario.save()

                    # -------------------------------------
                    # TIMESTAMP PRIMER COMENTARIO
                    # -------------------------------------

                    if es_primer_comentario:

                        ticket.primer_comentario_at = (
                            timezone.now()
                        )

                        ticket.save(
                            update_fields=[
                                "primer_comentario_at",
                                "actualizado_at",
                            ]
                        )

                        evento_comentario = (
                            HistorialTicket
                            .Evento
                            .PRIMER_COMENTARIO
                        )

                    else:

                        evento_comentario = (
                            HistorialTicket
                            .Evento
                            .COMENTARIO
                        )

                    # -------------------------------------
                    # HISTORIAL DEL COMENTARIO
                    # -------------------------------------

                    if (
                        comentario.tipo
                        == ComentarioTicket.Tipo.INTERNO
                    ):
                        descripcion_historial = (
                            "Se agregó una nota interna."
                        )
                    else:
                        descripcion_historial = (
                            "Se agregó un comentario "
                            "al ticket."
                        )

                    HistorialTicket.objects.create(
                        ticket=ticket,
                        usuario=request.user,
                        evento=evento_comentario,
                        origen=(
                            HistorialTicket
                            .Origen
                            .USUARIO
                        ),
                        descripcion=descripcion_historial,
                        valor_anterior={},
                        valor_nuevo={
                            "tipo": comentario.tipo,
                            "comentario_id": comentario.pk,
                            "comentario": (
                                comentario.comentario
                            ),
                        },
                    )

                    # -------------------------------------
                    # ARCHIVO OPCIONAL
                    # -------------------------------------

                    archivo_subido = (
                        archivo_form.cleaned_data
                        .get("archivo")
                    )

                    if archivo_subido:

                        archivo_ticket = (
                            archivo_form.save(
                                commit=False
                            )
                        )

                        archivo_ticket.ticket = ticket
                        archivo_ticket.comentario = comentario
                        archivo_ticket.usuario = request.user

                        archivo_ticket.nombre_original = (
                            archivo_subido.name
                        )

                        archivo_ticket.mime_type = (
                            getattr(
                                archivo_subido,
                                "content_type",
                                "",
                            )
                            or ""
                        )

                        archivo_ticket.size_bytes = (
                            archivo_subido.size
                        )

                        archivo_ticket.tipo = (
                            ArchivoTicket
                            .Tipo
                            .COMENTARIO
                        )

                        archivo_ticket.save()

                        HistorialTicket.objects.create(
                            ticket=ticket,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .ARCHIVO
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Se agregó un archivo "
                                "al comentario."
                            ),
                            valor_anterior={},
                            valor_nuevo={
                                "archivo": (
                                    archivo_ticket
                                    .nombre_original
                                ),
                                "mime_type": (
                                    archivo_ticket
                                    .mime_type
                                ),
                                "size_bytes": (
                                    archivo_ticket
                                    .size_bytes
                                ),
                            },
                        )

                messages.success(
                    request,
                    "Comentario agregado correctamente.",
                )

                return redirect(
                    "tickets:detalle",
                    folio=ticket.folio,
                )

            messages.error(
                request,
                (
                    "No se pudo guardar el comentario. "
                    "Revisa los campos indicados."
                ),
            )

        # =================================================
        # CIERRE CON COMENTARIO Y EVIDENCIA
        # =================================================

        elif accion == "cierre":
            cierre_form = CerrarTicketForm(
                {"comentario_cierre": request.POST.get("comentario", "")},
                {"evidencia_cierre": request.FILES.get("archivo")},
            )

            if cierre_form.is_valid():
                with transaction.atomic():
                    ticket = (
                        Ticket.objects
                        .select_for_update()
                        .get(pk=ticket.pk)
                    )

                    if ticket.estado_interno == EstadoTicket.CERRADO:
                        messages.info(request, "El ticket ya se encuentra cerrado.")
                        return redirect("tickets:detalle", folio=ticket.folio)

                    ahora = timezone.now()
                    estado_anterior = ticket.estado_interno
                    archivo_subido = cierre_form.cleaned_data["evidencia_cierre"]

                    comentario = ComentarioTicket.objects.create(
                        ticket=ticket,
                        usuario=request.user,
                        comentario=cierre_form.cleaned_data["comentario_cierre"],
                        tipo=ComentarioTicket.Tipo.CIERRE,
                    )

                    archivo_ticket = ArchivoTicket.objects.create(
                        ticket=ticket,
                        comentario=comentario,
                        usuario=request.user,
                        archivo=archivo_subido,
                        nombre_original=archivo_subido.name,
                        mime_type=(
                            getattr(archivo_subido, "content_type", "") or ""
                        ),
                        size_bytes=archivo_subido.size,
                        tipo=ArchivoTicket.Tipo.CIERRE,
                    )

                    ticket.estado_interno = EstadoTicket.CERRADO
                    ticket.cerrado_at = ahora
                    if ticket.resuelto_at is None:
                        ticket.resuelto_at = ahora
                    if ticket.primer_comentario_at is None:
                        ticket.primer_comentario_at = ahora
                    ticket.save(
                        update_fields=[
                            "estado_interno",
                            "cerrado_at",
                            "resuelto_at",
                            "primer_comentario_at",
                            "actualizado_at",
                        ]
                    )

                    sincronizar_sla(ticket, request.user, ahora)
                    HistorialTicket.objects.create(
                        ticket=ticket,
                        usuario=request.user,
                        evento=HistorialTicket.Evento.CERRADO,
                        origen=HistorialTicket.Origen.USUARIO,
                        descripcion=(
                            "El ticket fue cerrado con comentario y evidencia."
                        ),
                        valor_anterior={"estado": estado_anterior},
                        valor_nuevo={
                            "estado": EstadoTicket.CERRADO,
                            "comentario_id": comentario.pk,
                            "evidencia": archivo_ticket.nombre_original,
                        },
                    )

                messages.success(
                    request,
                    "Ticket cerrado correctamente con su evidencia.",
                )
                return redirect("tickets:detalle", folio=ticket.folio)

            messages.error(
                request,
                "Para cerrar el ticket debes escribir un comentario y adjuntar evidencia.",
            )
            comentario_form = ComentarioTicketForm(
                request.POST,
                permitir_nota_interna=puede_gestionar,
            )
            archivo_form = ArchivoTicketForm(request.POST, request.FILES)

        # =================================================
        # REAPERTURA EXCLUSIVA PARA ADMINISTRADORES
        # =================================================

        elif accion == "reapertura":
            reapertura_form = ReabrirTicketForm(request.POST)

            if reapertura_form.is_valid():
                with transaction.atomic():
                    ticket = (
                        Ticket.objects
                        .select_for_update()
                        .get(pk=ticket.pk)
                    )

                    if ticket.estado_interno != EstadoTicket.CERRADO:
                        messages.info(
                            request,
                            "Solo se puede reabrir un ticket cerrado.",
                        )
                        return redirect("tickets:detalle", folio=ticket.folio)

                    ahora = timezone.now()
                    motivo = reapertura_form.cleaned_data["motivo_reapertura"]

                    ComentarioTicket.objects.create(
                        ticket=ticket,
                        usuario=request.user,
                        comentario=motivo,
                        tipo=ComentarioTicket.Tipo.REAPERTURA,
                    )

                    ticket.estado_interno = EstadoTicket.EN_PROCESO
                    ticket.cerrado_at = None
                    ticket.resuelto_at = None
                    ticket.ultima_reapertura_at = ahora
                    ticket.numero_reaperturas += 1
                    ticket.save(
                        update_fields=[
                            "estado_interno",
                            "cerrado_at",
                            "resuelto_at",
                            "ultima_reapertura_at",
                            "numero_reaperturas",
                            "actualizado_at",
                        ]
                    )

                    sincronizar_sla(ticket, request.user, ahora)
                    HistorialTicket.objects.create(
                        ticket=ticket,
                        usuario=request.user,
                        evento=HistorialTicket.Evento.REABIERTO,
                        origen=HistorialTicket.Origen.USUARIO,
                        descripcion="El ticket fue reabierto por un administrador.",
                        valor_anterior={"estado": EstadoTicket.CERRADO},
                        valor_nuevo={
                            "estado": EstadoTicket.EN_PROCESO,
                            "motivo": motivo,
                        },
                    )

                messages.success(request, "Ticket reabierto correctamente.")
                return redirect("tickets:detalle", folio=ticket.folio)

            messages.error(request, "Escribe el motivo de la reapertura.")

        # =================================================
        # GESTIÓN COMPLETA
        # =================================================

        elif accion == "gestion":

            with transaction.atomic():

                ticket = (
                    Ticket.objects
                    .select_for_update()
                    .get(pk=ticket.pk)
                )

                # -----------------------------------------
                # Snapshot ANTES de validar el ModelForm.
                #
                # Importante:
                # ModelForm modifica instance durante
                # is_valid(), por eso guardamos esto antes.
                # -----------------------------------------

                anterior = {
                    "estado_interno": ticket.estado_interno,
                    "prioridad": ticket.prioridad,
                    "responsable": ticket.responsable,
                    "partner": ticket.partner,
                    "grupo": ticket.grupo,
                    "subgrupo": ticket.subgrupo,
                    "dependencia_actual": (
                        ticket.dependencia_actual
                    ),
                    "estatus_operativo_actual": (
                        ticket.estatus_operativo_actual
                    ),
                }

                gestion_form = GestionTicketForm(
                    request.POST,
                    instance=ticket,
                )

                if gestion_form.is_valid():

                    ticket_actualizado = (
                        gestion_form.save(
                            commit=False
                        )
                    )

                    ahora = timezone.now()

                    # =====================================
                    # REGLA AUTOMÁTICA
                    # NUEVO + RESPONSABLE -> ASIGNADO
                    # =====================================

                    if (
                        ticket_actualizado.responsable
                        and ticket_actualizado.estado_interno
                        == EstadoTicket.NUEVO
                    ):
                        ticket_actualizado.estado_interno = (
                            EstadoTicket.ASIGNADO
                        )

                    estado_nuevo = (
                        ticket_actualizado.estado_interno
                    )

                    estado_anterior = (
                        anterior["estado_interno"]
                    )

                    responsable_nuevo = (
                        ticket_actualizado.responsable
                    )

                    responsable_anterior = (
                        anterior["responsable"]
                    )

                    partner_nuevo = ticket_actualizado.partner
                    partner_anterior = anterior["partner"]

                    # =====================================
                    # TIMESTAMP RESPONSABLE
                    # =====================================

                    if (
                        responsable_anterior
                        != responsable_nuevo
                    ):

                        if responsable_nuevo:
                            ticket_actualizado.asignado_at = (
                                ahora
                            )
                        else:
                            ticket_actualizado.asignado_at = (
                                None
                            )

                    # =====================================
                    # REAPERTURA
                    # =====================================

                    estados_finalizados = {
                        EstadoTicket.RESUELTO,
                        EstadoTicket.CERRADO,
                    }

                    es_reapertura = (
                        estado_anterior
                        in estados_finalizados
                        and estado_nuevo
                        not in estados_finalizados
                    )

                    if es_reapertura:

                        ticket_actualizado.ultima_reapertura_at = (
                            ahora
                        )
                        ticket_actualizado.numero_reaperturas = (
                            ticket.numero_reaperturas + 1
                        )

                        # Ya no está cerrado/resuelto
                        # actualmente.
                        ticket_actualizado.resuelto_at = None
                        ticket_actualizado.cerrado_at = None

                    # =====================================
                    # RESUELTO
                    # =====================================

                    elif (
                        estado_nuevo
                        == EstadoTicket.RESUELTO
                        and estado_anterior
                        != EstadoTicket.RESUELTO
                    ):

                        ticket_actualizado.resuelto_at = (
                            ahora
                        )

                        ticket_actualizado.cerrado_at = None

                    # =====================================
                    # CERRADO
                    # =====================================

                    elif (
                        estado_nuevo
                        == EstadoTicket.CERRADO
                        and estado_anterior
                        != EstadoTicket.CERRADO
                    ):

                        ticket_actualizado.cerrado_at = (
                            ahora
                        )

                    # =====================================
                    # GUARDAR TICKET
                    # =====================================

                    ticket_actualizado.save()
                    sincronizar_sla(ticket_actualizado, request.user, ahora)
                    if (anterior["dependencia_actual"] != ticket_actualizado.dependencia_actual or anterior["estatus_operativo_actual"] != ticket_actualizado.estatus_operativo_actual):
                        registrar_segmento_operativo(ticket_actualizado, ahora)

                    # =====================================
                    # HISTORIAL · PARTNER
                    # =====================================

                    if partner_anterior != partner_nuevo:
                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=HistorialTicket.Evento.PARTNER,
                            origen=HistorialTicket.Origen.USUARIO,
                            descripcion="Cambio de partner de soporte.",
                            valor_anterior={
                                "partner": usuario_valor(partner_anterior)
                            },
                            valor_nuevo={
                                "partner": usuario_valor(partner_nuevo)
                            },
                        )

                    # =====================================
                    # HISTORIAL · RESPONSABLE
                    # =====================================

                    if (
                        responsable_anterior
                        != responsable_nuevo
                    ):

                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .ASIGNADO
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Cambio de responsable "
                                "CAROBRA."
                            ),
                            valor_anterior={
                                "responsable": usuario_valor(
                                    responsable_anterior
                                )
                            },
                            valor_nuevo={
                                "responsable": usuario_valor(
                                    responsable_nuevo
                                )
                            },
                        )

                    # =====================================
                    # HISTORIAL · ESTADO
                    # =====================================

                    if estado_anterior != estado_nuevo:

                        if es_reapertura:

                            evento_estado = (
                                HistorialTicket
                                .Evento
                                .REABIERTO
                            )

                            descripcion_estado = (
                                "El ticket fue reabierto."
                            )

                        elif (
                            estado_nuevo
                            == EstadoTicket.RESUELTO
                        ):

                            evento_estado = (
                                HistorialTicket
                                .Evento
                                .RESUELTO
                            )

                            descripcion_estado = (
                                "El ticket fue marcado "
                                "como resuelto."
                            )

                        elif (
                            estado_nuevo
                            == EstadoTicket.CERRADO
                        ):

                            evento_estado = (
                                HistorialTicket
                                .Evento
                                .CERRADO
                            )

                            descripcion_estado = (
                                "El ticket fue cerrado."
                            )

                        else:

                            evento_estado = (
                                HistorialTicket
                                .Evento
                                .ESTADO
                            )

                            if (
                                estado_anterior
                                == EstadoTicket.NUEVO
                                and estado_nuevo
                                == EstadoTicket.ASIGNADO
                                and responsable_nuevo
                            ):
                                descripcion_estado = (
                                    "El ticket cambió "
                                    "automáticamente de "
                                    "Nuevo a Asignado al "
                                    "recibir un responsable."
                                )
                            else:
                                descripcion_estado = (
                                    "Cambio de estado "
                                    "del ticket."
                                )

                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=evento_estado,
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=descripcion_estado,
                            valor_anterior={
                                "estado": estado_anterior
                            },
                            valor_nuevo={
                                "estado": estado_nuevo
                            },
                        )

                    # =====================================
                    # HISTORIAL · PRIORIDAD
                    # =====================================

                    if (
                        anterior["prioridad"]
                        != ticket_actualizado.prioridad
                    ):

                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .PRIORIDAD
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Cambio de prioridad "
                                "del ticket."
                            ),
                            valor_anterior={
                                "prioridad": (
                                    anterior["prioridad"]
                                )
                            },
                            valor_nuevo={
                                "prioridad": (
                                    ticket_actualizado
                                    .prioridad
                                )
                            },
                        )

                    # =====================================
                    # HISTORIAL · GRUPO
                    # =====================================

                    if (
                        anterior["grupo"]
                        != ticket_actualizado.grupo
                    ):

                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .GRUPO
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Cambio de grupo "
                                "de soporte."
                            ),
                            valor_anterior={
                                "grupo": objeto_valor(
                                    anterior["grupo"]
                                )
                            },
                            valor_nuevo={
                                "grupo": objeto_valor(
                                    ticket_actualizado.grupo
                                )
                            },
                        )

                    # =====================================
                    # HISTORIAL · SUBGRUPO
                    # =====================================

                    if (
                        anterior["subgrupo"]
                        != ticket_actualizado.subgrupo
                    ):

                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .SUBGRUPO
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Cambio de subgrupo "
                                "de soporte."
                            ),
                            valor_anterior={
                                "subgrupo": objeto_valor(
                                    anterior["subgrupo"]
                                )
                            },
                            valor_nuevo={
                                "subgrupo": objeto_valor(
                                    ticket_actualizado
                                    .subgrupo
                                )
                            },
                        )

                    # =====================================
                    # HISTORIAL · DEPENDENCIA
                    # =====================================

                    if (
                        anterior["dependencia_actual"]
                        != ticket_actualizado
                        .dependencia_actual
                    ):

                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .DEPENDENCIA
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Cambio de dependencia "
                                "del ticket."
                            ),
                            valor_anterior={
                                "dependencia": objeto_valor(
                                    anterior[
                                        "dependencia_actual"
                                    ]
                                )
                            },
                            valor_nuevo={
                                "dependencia": objeto_valor(
                                    ticket_actualizado
                                    .dependencia_actual
                                )
                            },
                        )

                    # =====================================
                    # HISTORIAL · ESTATUS OPERATIVO
                    # =====================================

                    if (
                        anterior[
                            "estatus_operativo_actual"
                        ]
                        != ticket_actualizado
                        .estatus_operativo_actual
                    ):

                        HistorialTicket.objects.create(
                            ticket=ticket_actualizado,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .ESTATUS_OPERATIVO
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Cambio de estatus "
                                "operativo."
                            ),
                            valor_anterior={
                                "estatus_operativo": (
                                    objeto_valor(
                                        anterior[
                                            "estatus_operativo_actual"
                                        ]
                                    )
                                )
                            },
                            valor_nuevo={
                                "estatus_operativo": (
                                    objeto_valor(
                                        ticket_actualizado
                                        .estatus_operativo_actual
                                    )
                                )
                            },
                        )

                    messages.success(
                        request,
                        "Ticket actualizado correctamente.",
                    )

                    return redirect(
                        "tickets:detalle",
                        folio=ticket_actualizado.folio,
                    )

                messages.error(
                    request,
                    (
                        "No se pudieron guardar los cambios. "
                        "Revisa los campos indicados."
                    ),
                )

        # =================================================
        # ASIGNACIÓN RÁPIDA ACTUAL
        # =================================================

        else:

            form_responsable = AsignarResponsableForm(
                request.POST
            )

            if form_responsable.is_valid():

                nuevo_responsable = (
                    form_responsable.cleaned_data[
                        "responsable"
                    ]
                )

                with transaction.atomic():

                    ticket = (
                        Ticket.objects
                        .select_for_update()
                        .get(pk=ticket.pk)
                    )

                    responsable_anterior = (
                        ticket.responsable
                    )

                    estado_anterior = (
                        ticket.estado_interno
                    )

                    if (
                        responsable_anterior
                        != nuevo_responsable
                    ):

                        ticket.responsable = (
                            nuevo_responsable
                        )

                        if nuevo_responsable:
                            ticket.asignado_at = (
                                timezone.now()
                            )
                        else:
                            ticket.asignado_at = None

                        cambio_automatico_estado = (
                            nuevo_responsable is not None
                            and ticket.estado_interno
                            == EstadoTicket.NUEVO
                        )

                        if cambio_automatico_estado:
                            ticket.estado_interno = (
                                EstadoTicket.ASIGNADO
                            )

                        ticket.save()

                        HistorialTicket.objects.create(
                            ticket=ticket,
                            usuario=request.user,
                            evento=(
                                HistorialTicket
                                .Evento
                                .ASIGNADO
                            ),
                            origen=(
                                HistorialTicket
                                .Origen
                                .USUARIO
                            ),
                            descripcion=(
                                "Cambio de responsable "
                                "CAROBRA."
                            ),
                            valor_anterior={
                                "responsable": (
                                    usuario_valor(
                                        responsable_anterior
                                    )
                                )
                            },
                            valor_nuevo={
                                "responsable": (
                                    usuario_valor(
                                        nuevo_responsable
                                    )
                                )
                            },
                        )

                        if cambio_automatico_estado:

                            HistorialTicket.objects.create(
                                ticket=ticket,
                                usuario=request.user,
                                evento=(
                                    HistorialTicket
                                    .Evento
                                    .ESTADO
                                ),
                                origen=(
                                    HistorialTicket
                                    .Origen
                                    .SISTEMA
                                ),
                                descripcion=(
                                    "El ticket cambió "
                                    "automáticamente de "
                                    "Nuevo a Asignado al "
                                    "recibir un responsable."
                                ),
                                valor_anterior={
                                    "estado": estado_anterior
                                },
                                valor_nuevo={
                                    "estado": (
                                        EstadoTicket
                                        .ASIGNADO
                                    )
                                },
                            )

                        messages.success(
                            request,
                            (
                                "Responsable actualizado "
                                "correctamente."
                            ),
                        )

                    else:

                        messages.info(
                            request,
                            "El responsable no cambió.",
                        )

                return redirect(
                    "tickets:detalle",
                    folio=ticket.folio,
                )

    # =====================================================
    # FORMULARIOS PARA RENDER
    # =====================================================

    form_responsable = AsignarResponsableForm(
        initial={
            "responsable": ticket.responsable
        }
    )

    # Si hubo un POST de gestión inválido, conservamos
    # ese formulario con sus errores.
    if not (
        request.method == "POST"
        and request.POST.get("accion") == "gestion"
    ):
        gestion_form = GestionTicketForm(
            instance=ticket
        )


    # =====================================================
    # COMENTARIOS RECIENTES
    # =====================================================

    comentarios_query = (
        ticket.comentarios
        .select_related("usuario")
        .prefetch_related("archivos")
    )

    if not puede_gestionar:
        comentarios_query = comentarios_query.exclude(
            tipo=ComentarioTicket.Tipo.INTERNO
        )

    comentarios = comentarios_query.order_by("-creado_at")[:50]

    # =====================================================
    # HISTORIAL RECIENTE
    # =====================================================

    historial_query = (
        ticket.historial
        .select_related("usuario")
    )

    if not puede_gestionar:
        historial_query = historial_query.exclude(
            valor_nuevo__tipo=ComentarioTicket.Tipo.INTERNO
        )

    historial = historial_query.order_by("-creado_at")[:30]

    sla = resumen_sla(ticket)
    if sla["limite"]:
        sla["porcentaje"] = min(
            100,
            round((sla["efectivo"] / sla["limite"]) * 100),
        )
    else:
        sla["porcentaje"] = 0
    sla["en_riesgo"] = bool(
        sla["limite"]
        and not sla["excedido"]
        and sla["efectivo"] >= sla["limite"] * 0.75
    )

    return render(
        request,
        "tickets/detalle_ticket.html",
        {
            "ticket": ticket,
            "activos_tienda": activos_tienda,
            "form_responsable": form_responsable,
            "gestion_form": gestion_form,
            "comentario_form": comentario_form,
            "archivo_form": archivo_form,
            "cierre_form": cierre_form,
            "reapertura_form": reapertura_form,
            "comentarios": comentarios,
            "historial": historial,
            "puede_gestionar": puede_gestionar,
            "rol_en_ticket": rol_en_ticket,
            "es_responsable": es_responsable,
            "es_partner": es_partner,
            "es_creador": es_creador,
            "partner_usuario": partner_usuario,
            "sla": sla,
            "cobertura_actual": seguimiento_actual(ticket),
            "puede_seguimiento": participa_en_ticket(request.user, ticket) and permiso_ticket(request.user, "puede_seguimiento"),
            "puede_cerrar": participa_en_ticket(request.user, ticket) and permiso_ticket(request.user, "puede_cerrar"),
            "puede_reasignar": participa_en_ticket(request.user, ticket) and permiso_ticket(request.user, "puede_reasignar"),
            "seguimiento_form": SeguimientoForm(initial={"estado": ticket.estado_interno, "siguiente_accion": ticket.siguiente_accion, "motivo_espera": ticket.motivo_espera, "esperando_a": ticket.esperando_a}),
            "suplencia_form": SuplenciaForm() if puede_gestionar else None,
            "reglas_clasificacion": ConfiguracionSLA.objects.filter(activo=True) if puede_gestionar else [],
        },
    )


@login_required
def descargar_archivo_ticket(request, archivo_id):
    archivo = get_object_or_404(
        ArchivoTicket.objects.select_related("ticket", "comentario"),
        pk=archivo_id,
    )

    ticket_visible = tickets_visibles_para(
        request.user,
        Ticket.objects.filter(pk=archivo.ticket_id),
    ).exists()

    es_nota_interna = (
        archivo.comentario_id
        and archivo.comentario.tipo == ComentarioTicket.Tipo.INTERNO
    )

    if not ticket_visible or (
        es_nota_interna and not es_administrador(request.user)
    ):
        raise Http404("El archivo no está disponible.")

    if not archivo.archivo:
        raise Http404("El archivo no está disponible.")

    try:
        contenido = archivo.archivo.open("rb")
    except (FileNotFoundError, OSError):
        raise Http404("El archivo no está disponible.")

    nombre = archivo.nombre_original or Path(archivo.archivo.name).name
    content_type = mimetypes.guess_type(nombre)[0] or "application/octet-stream"
    tipos_inline = {
        "application/pdf",
        "image/jpeg",
        "image/png",
        "image/webp",
    }
    respuesta = FileResponse(
        contenido,
        as_attachment=content_type not in tipos_inline,
        filename=nombre,
        content_type=content_type,
    )
    respuesta["Cache-Control"] = "private, no-store"
    respuesta["X-Content-Type-Options"] = "nosniff"
    return respuesta
